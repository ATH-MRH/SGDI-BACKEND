"""Import Excel de candidats (recrute.irongs.com) : Excel → vérification → confirmation → fiches."""
import io
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.security import hash_password
from app.modules.auth.models import AuditEvent, User
from app.modules.drh import candidate_import as importer
from app.modules.drh.candidate_import_models import CandidateImportSession
from app.modules.drh.candidate_import_reference import FIELDS, WILAYAS, field_for_header, norm_key
from app.modules.drh.models import Candidate, Employee
from app.modules.irongs.models import Position
from tests.module_cleanup import purge_rows_created_by_this_module  # noqa: F401 (fixture autouse)

BASE = "/api/drh/candidates/import"
SECURITE, SOLUTION = "IRON GLOBAL SECURITE", "IRON GLOBAL SOLUTION"
SCOPE_SECURITE, SCOPE_SOLUTION = "Iron Global Securite", "Iron Global Solution"
HEADERS = ["Nom", "Prénom", "Date de naissance", "Téléphone", "E-mail", "Wilaya", "Commune", "Poste souhaité", "Sexe"]


def _tag() -> str:
    return uuid.uuid4().hex[:8].upper()


def _phone() -> str:
    return f"05{uuid.uuid4().int % 10**8:08d}"


def _xlsx(rows, headers=HEADERS, *, sheet="Candidats", extra_sheets=None) -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for name, content in (extra_sheets or {}).items():
        other = book.create_sheet(name)
        for row in content:
            other.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _xls(rows, headers=HEADERS) -> bytes:
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    ws = book.add_sheet("Candidats")
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, value in enumerate(headers):
        ws.write(0, c, value)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            if isinstance(value, (date, datetime)):
                ws.write(r, c, value, date_style)
            elif value is not None:
                ws.write(r, c, value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _upload(client, headers, content, name="candidats.xlsx"):
    return client.post(BASE, headers=headers, files={"file": (name, content, "application/octet-stream")})


def _preview(client, headers, upload, mapping=None, sheet=None):
    body = upload.json()
    sheet = body["default_sheet"] if sheet is None else sheet
    columns = body["sheets"][sheet]["columns"]
    mapping = mapping if mapping is not None else {str(column["index"]): column["suggested"] for column in columns}
    return client.post(f"{BASE}/{body['session_id']}/preview", headers=headers, json={"sheet": sheet, "mapping": mapping})


def _confirm(client, headers, session_id, decisions=None):
    return client.post(f"{BASE}/{session_id}/confirm", headers=headers, json={"decisions": decisions or {}})


def _run(client, headers, content, *, decisions=None, name="candidats.xlsx", mapping=None):
    upload = _upload(client, headers, content, name)
    assert upload.status_code == 200, upload.text
    preview = _preview(client, headers, upload, mapping)
    assert preview.status_code == 200, preview.text
    confirm = _confirm(client, headers, upload.json()["session_id"], decisions)
    assert confirm.status_code == 200, confirm.text
    return preview.json(), confirm.json()


def _row(tag, **overrides):
    values = {"Nom": f"IMPORT{tag}", "Prénom": "Karim", "Date de naissance": "15/03/1990", "Téléphone": _phone(),
              "E-mail": f"imp{tag.lower()}@example.com", "Wilaya": "Alger", "Commune": "Alger Centre",
              "Poste souhaité": "Chauffeur", "Sexe": "M", **overrides}
    return [values[header] for header in HEADERS]


def _by_name(db, tag):
    db.expire_all()
    return db.execute(select(Candidate).where(Candidate.last_name == f"IMPORT{tag}")).scalars().all()


def _recruiter(client, db, *, societies=(SCOPE_SECURITE,), actions=("read", "create", "update")):
    username = f"REC{_tag()}"
    db.add(User(username=username, email=f"{username.lower()}@test.com", full_name="Recruteur Import", role="recruteur",
                access_level="H3", authorized_societies=list(societies), authorized_structures=[], authorized_modules=["recrute"],
                authorized_actions=list(actions), password_hash=hash_password("recruteur-pass-123"), is_active=True))
    db.commit()
    login = client.post("/api/auth/login", json={"username": username, "password": "recruteur-pass-123"})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


@pytest.fixture(scope="module", autouse=True)
def _purge_candidates_created_by_this_module():
    """Base de test partagée : les candidats et sessions créés ici fausseraient les tests
    suivants qui lisent des listes ou des compteurs par société."""
    from sqlalchemy import delete, func
    from tests.conftest import TestSessionLocal

    with TestSessionLocal() as session:
        mark = session.scalar(select(func.max(Candidate.id))) or 0
    yield
    with TestSessionLocal() as session:
        session.execute(delete(Candidate).where(Candidate.id > mark))
        session.execute(delete(CandidateImportSession))
        session.commit()


@pytest.fixture(autouse=True)
def _positions(db):
    if not db.query(Position).filter(Position.name == "Chauffeur").first():
        db.add(Position(name="Chauffeur", society=None))
        db.commit()


# ── Lecture, correspondance, création ────────────────────────────────────────────────────
def test_xlsx_import_creates_candidates_in_the_initial_status(client, auth_headers, db):
    tag = _tag()
    row = _row(tag, **{"Téléphone": "0555 12 34 56".replace("12 34 56", f"{uuid.uuid4().int % 10**6:06d}")})
    preview, result = _run(client, auth_headers, _xlsx([row]))
    assert preview["counts"] == {"analyzed": 1, "valid": 1, "duplicates": 0, "errors": 0, "warnings": 0, "empty_rows": 0}
    assert result["counts"]["created"] == 1 and result["counts"]["failed"] == 0 and result["already_processed"] is False
    [candidate] = _by_name(db, tag)
    assert candidate.status == "nouvelle" and candidate.society is None             # vivier Groupe, statut initial
    assert candidate.first_name == "KARIM" and candidate.desired_position == "CHAUFFEUR"
    assert candidate.phone.startswith("0555") and len(candidate.phone) == 10 and isinstance(candidate.phone, str)
    data = candidate.data
    assert data["dateNaissance"] == "1990-03-15" and data["wilaya"] == "Alger" and data["commune"] == "Alger Centre"
    assert data["sexe"] == "M" and data["posteSouhaite"] == "Chauffeur" and data["telephone"] == candidate.phone
    assert data["importSource"] == "candidats.xlsx" and data["importRow"] == 2 and data["moduleOrigine"] == "recrutement"
    assert data["auditTrail"][0]["action"] == "creation"                            # service métier existant
    for forbidden in ("avisDecision", "fichePositionValidee", "sectionValidations", "allowDuplicate"):
        assert forbidden not in data
    # La fiche apparaît dans « Nouvelles candidatures » de recrute.irongs.com.
    page = client.get("/api/drh/candidates/page", headers=auth_headers, params={"mode": "new", "q": f"IMPORT{tag}"})
    assert [item["id"] for item in page.json()["items"]] == [candidate.id]


def test_xls_import_reads_dates_and_numeric_phones(client, auth_headers, db):
    tag = _tag()
    number = int(f"5{uuid.uuid4().int % 10**8:08d}")                                # zéro initial perdu par Excel
    row = _row(tag, **{"Date de naissance": date(1988, 12, 25), "Téléphone": number})
    preview, result = _run(client, auth_headers, _xls([row]), name="candidats.xls")
    assert result["counts"]["created"] == 1
    [candidate] = _by_name(db, tag)
    assert candidate.phone == f"0{number}" and candidate.data["dateNaissance"] == "1988-12-25"
    [warning] = preview["rows"][0]["warnings"]
    assert warning["field"] == "phone" and "Zéro initial" in warning["message"]


def test_header_variants_are_recognised_and_the_mapping_can_be_corrected(client, auth_headers, db):
    for header, expected in (("NOM DE FAMILLE", "last_name"), ("  prénoms ", "first_name"), ("Tél.", "phone"), ("E-MAIL", "email"),
                             ("Né(e) le", "birth_date"), ("Poste recherché", "desired_position"), ("Situation familiale", "family_status"),
                             ("Observations", "notes"), ("Source de candidature", "source"), ("اللقب", "last_name"), ("الاسم", "first_name"),
                             ("تاريخ الميلاد", "birth_date"), ("رقم الهاتف", "phone"), ("الولاية", "wilaya"), ("Diplôme", None)):
        assert field_for_header(header) == expected, header
    tag = _tag()
    headers = ["Patronyme", "الاسم", "GSM", "Niveau d'études", "Nom"]
    upload = _upload(client, auth_headers, _xlsx([[f"IMPORT{tag}", "Samir", _phone(), "Licence", "ALIAS"]], headers))
    columns = upload.json()["sheets"][0]["columns"]
    assert [column["suggested"] for column in columns] == [None, "first_name", "phone", None, "last_name"]
    assert columns[0]["samples"] == [f"IMPORT{tag}"] and columns[0]["letter"] == "A"
    session_id = upload.json()["session_id"]
    # Champ associé à deux colonnes → correspondance contradictoire refusée.
    conflict = client.post(f"{BASE}/{session_id}/preview", headers=auth_headers,
                           json={"sheet": 0, "mapping": {"0": "last_name", "1": "first_name", "4": "last_name"}})
    assert conflict.status_code == 422 and "contradictoire" in conflict.json()["detail"]
    missing = client.post(f"{BASE}/{session_id}/preview", headers=auth_headers, json={"sheet": 0, "mapping": {"1": "first_name"}})
    assert missing.status_code == 422 and "Nom" in missing.json()["detail"]
    unknown = client.post(f"{BASE}/{session_id}/preview", headers=auth_headers, json={"sheet": 0, "mapping": {"0": "salaire_secret"}})
    assert unknown.status_code == 422
    # Correction manuelle : A = Nom ; la colonne « Nom » d'origine et « Niveau d'études » sont ignorées.
    preview = client.post(f"{BASE}/{session_id}/preview", headers=auth_headers,
                          json={"sheet": 0, "mapping": {"0": "last_name", "1": "first_name", "2": "phone", "3": None, "4": "__ignore__"}})
    assert preview.status_code == 200, preview.text
    assert len(preview.json()["ignored_columns"]) == 2
    assert _confirm(client, auth_headers, session_id).json()["counts"]["created"] == 1
    [candidate] = _by_name(db, tag)
    assert candidate.first_name == "SAMIR" and "niveauEtude" not in candidate.data      # aucun champ inventé


def test_sheet_choice_and_template_helper_sheets(client, auth_headers, db):
    tag = _tag()
    content = _xlsx([["x", "y"]], ["Divers", "Autre"], sheet="Notes", extra_sheets={
        "Avril": [HEADERS, _row(tag)], "Exemple": [["EXEMPLE — ne pas importer cette feuille"], [], HEADERS, _row(_tag())]})
    upload = _upload(client, auth_headers, content)
    body = upload.json()
    assert [(sheet["name"], sheet["importable"]) for sheet in body["sheets"]] == [("Notes", True), ("Avril", True), ("Exemple", False)]
    assert body["default_sheet"] == 1                                              # feuille aux en-têtes reconnus
    refused = _preview(client, auth_headers, upload, sheet=2, mapping={})
    assert refused.status_code == 422 and "non importable" in refused.json()["detail"]
    assert _preview(client, auth_headers, upload).status_code == 200
    assert _confirm(client, auth_headers, body["session_id"]).json()["counts"]["created"] == 1
    assert len(_by_name(db, tag)) == 1


# ── Validation ───────────────────────────────────────────────────────────────────────────
def test_dates_phones_references_and_required_fields_are_validated(client, auth_headers, db):
    tags = [_tag() for _ in range(12)]
    rows = [
        _row(tags[0], **{"Date de naissance": "03/04/1991"}),                       # ambiguë → avertissement
        _row(tags[1], **{"Date de naissance": 33000}),                              # numéro de série Excel
        _row(tags[2], **{"Date de naissance": "04/25/1990"}),                       # mois/jour → erreur
        _row(tags[3], **{"Date de naissance": "15/03/90"}),                         # année sur 2 chiffres
        _row(tags[4], **{"Téléphone": "+213 661 12 34 56".replace("12 34 56", f"{uuid.uuid4().int % 10**6:06d}")}),
        _row(tags[5], **{"Téléphone": "12345"}),
        _row(tags[6], **{"E-mail": "pas-un-email"}),
        _row(tags[7], **{"Poste souhaité": "Astronaute"}),
        _row(tags[8], **{"Wilaya": "Atlantide"}),
        _row(tags[9], **{"Wilaya": "Oran", "Commune": "Alger Centre"}),             # commune hors wilaya
        _row(tags[10], **{"Prénom": None}),                                         # obligatoire absent
        _row(tags[11], **{"Date de naissance": date.today() - timedelta(days=365 * 15)}),   # règle serveur : 20 ans
        [None] * len(HEADERS),                                                      # ligne vide ignorée
        _row(_tag(), **{"Sexe": "X"}),
        _row(_tag(), **{"Date de naissance": "=TODAY()"}),                          # formule sans valeur enregistrée
    ]
    positions_before = db.query(Position).count()
    upload = _upload(client, auth_headers, _xlsx(rows))
    preview = _preview(client, auth_headers, upload).json()
    by_row = {row["row"]: row for row in preview["rows"]}
    assert preview["counts"]["analyzed"] == 14 and preview["counts"]["empty_rows"] == 1 and 14 not in by_row

    def error(row_number):
        assert by_row[row_number]["status"] == "error", by_row[row_number]
        return by_row[row_number]["errors"][0]

    assert by_row[2]["status"] == "valid" and "ambiguë" in by_row[2]["warnings"][0]["message"]
    assert by_row[2]["values"]["birth_date"] == "03/04/1991"
    assert by_row[3]["status"] == "valid" and "Nombre Excel 33000" in by_row[3]["warnings"][0]["message"]
    assert error(4)["field"] == "birth_date" and "mois/jour" in error(4)["message"]
    assert "deux chiffres" in error(5)["message"]
    assert by_row[6]["status"] == "valid" and by_row[6]["values"]["phone"].startswith("0661") and not by_row[6]["warnings"]
    assert error(7)["field"] == "phone"
    assert error(8)["field"] == "email"
    assert error(9)["field"] == "desired_position" and "référentiel" in error(9)["message"]
    assert error(10)["field"] == "wilaya"
    assert error(11)["field"] == "commune" and "Oran" in error(11)["message"]
    assert error(12) == {"field": "first_name", "label": "Prénom", "message": "Champ obligatoire absent"}
    assert error(13)["field"] == "birth_date" and "20 ans" in error(13)["message"]      # validation serveur du formulaire
    assert error(15)["field"] == "sex"
    assert error(16)["field"] == "birth_date" and "formule" in error(16)["message"]
    result = _confirm(client, auth_headers, upload.json()["session_id"]).json()
    assert result["counts"] == {"analyzed": 14, "created": 3, "updated": 0, "skipped": 0, "errors": 11, "failed": 0, "empty_rows": 1}
    assert [len(_by_name(db, tag)) for tag in tags] == [1, 1, 0, 0, 1, 0, 0, 0, 0, 0, 0, 0]
    assert _by_name(db, tags[1])[0].data["dateNaissance"] == (date(1899, 12, 30) + timedelta(days=33000)).isoformat()
    assert db.query(Position).count() == positions_before                           # aucune référence créée
    assert {row["row"] for row in result["rows"] if row["outcome"] == "error"} == {4, 5, 7, 8, 9, 10, 11, 12, 13, 15, 16}


def test_formula_cells_use_only_a_cached_value(db):
    """openpyxl n'évalue rien : une formule sans valeur enregistrée est signalée, jamais calculée."""
    book = Workbook()
    ws = book.active
    ws.append(["Nom", "Prénom"])
    ws.append(["=CONCATENATE(\"A\",\"B\")", "Lina"])
    buffer = io.BytesIO()
    book.save(buffer)
    workbook = importer.read_workbook("f.xlsx", buffer.getvalue())
    assert workbook["sheets"][0]["rows"][1][0] == {"e": "formule sans valeur enregistrée"}
    # Même classeur, avec la valeur calculée enregistrée par Excel (<v>) : elle seule est lue.
    patched = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = raw.decode()
                assert "<v></v>" in text or "<v/>" in text or "<v />" in text, text
                raw = text.replace("<v></v>", "<v>AB</v>").replace("<v/>", "<v>AB</v>").replace("<v />", "<v>AB</v>").replace('<c r="A2"', '<c r="A2" t="str"').encode()
            target.writestr(item, raw)
    cached = importer.read_workbook("f.xlsx", patched.getvalue())
    assert cached["sheets"][0]["rows"][1][0] == "AB"


# ── Doublons ─────────────────────────────────────────────────────────────────────────────
def test_duplicates_inside_the_file_and_with_existing_candidates_are_skipped_by_default(client, auth_headers, db):
    tag, other, third = _tag(), _tag(), _tag()
    phone, email = _phone(), f"dup{tag.lower()}@example.com"
    existing = client.post("/api/drh/candidates", headers=auth_headers, json={
        "first_name": "Karim", "last_name": f"IMPORT{tag}", "phone": phone, "email": email, "status": "nouvelle",
        "data": {"dateNaissance": "1990-03-15"}}).json()["data"]
    shared_phone = _phone()
    rows = [
        _row(tag, **{"Téléphone": "+213" + phone[1:], "E-mail": f"autre{tag.lower()}@example.com", "Date de naissance": "01/01/1980"}),  # téléphone
        _row(_tag(), **{"E-mail": email.upper()}),                                  # e-mail (casse ignorée)
        _row(tag, **{"Téléphone": _phone(), "E-mail": None}),                       # nom + prénom + naissance
        _row(other, **{"Téléphone": shared_phone}),
        _row(third, **{"Téléphone": shared_phone}),                                 # doublon de la ligne 5
    ]
    preview, result = _run(client, auth_headers, _xlsx(rows))
    by_row = {row["row"]: row for row in preview["rows"]}
    assert [by_row[n]["status"] for n in (2, 3, 4, 5, 6)] == ["duplicate", "duplicate", "duplicate", "valid", "duplicate"]
    assert by_row[2]["matches"] == [{"type": "candidat", "reasons": ["telephone"], "visible": True, "updatable": True, "id": existing["id"],
                                    "name": f"IMPORT{tag} KARIM", "society": None, "state": "Nouvelle candidature"}]
    assert by_row[3]["matches"][0]["reasons"] == ["email"] and by_row[4]["matches"][0]["reasons"] == ["identite"]
    assert by_row[6]["internal"] == [{"row": 5, "reasons": ["telephone"]}] and by_row[6]["matches"] == []
    assert all(by_row[n]["default_action"] == "skip" for n in (2, 3, 4, 6))
    assert result["counts"] == {"analyzed": 5, "created": 1, "updated": 0, "skipped": 4, "errors": 0, "failed": 0, "empty_rows": 0}
    assert len(_by_name(db, tag)) == 1 and len(_by_name(db, other)) == 1 and _by_name(db, third) == []
    skipped = {row["row"]: row["reason"] for row in result["rows"] if row["outcome"] == "skipped"}
    assert f"dossier n° {existing['id']}" in skipped[2] and "ligne 5 du fichier" in skipped[6]


def test_explicit_creation_despite_a_match_and_no_automatic_merge_when_ambiguous(client, auth_headers, db):
    tag_a, tag_b, tag_c = _tag(), _tag(), _tag()
    phone, email = _phone(), f"amb{tag_a.lower()}@example.com"
    first = client.post("/api/drh/candidates", headers=auth_headers, json={"first_name": "Ali", "last_name": f"IMPORT{tag_a}", "phone": phone, "status": "nouvelle"}).json()["data"]
    second = client.post("/api/drh/candidates", headers=auth_headers, json={"first_name": "Ali", "last_name": f"IMPORT{tag_b}", "email": email, "status": "nouvelle"}).json()["data"]
    content = _xlsx([_row(tag_c, **{"Téléphone": phone, "E-mail": email})])      # deux fiches différentes correspondent
    upload = _upload(client, auth_headers, content)
    row = _preview(client, auth_headers, upload).json()["rows"][0]
    assert {match["id"] for match in row["matches"]} == {first["id"], second["id"]} and row["ambiguous"] is True
    assert row["actions"] == ["skip", "create"] and row["update_target"] is None    # jamais de fusion proposée
    forced_update = _confirm(client, auth_headers, upload.json()["session_id"], {"2": "update"}).json()
    assert forced_update["counts"]["failed"] == 1 and forced_update["counts"]["updated"] == 0
    assert _by_name(db, tag_c) == []
    # Création malgré correspondance : choix explicite uniquement.
    _, result = _run(client, auth_headers, content, decisions={"2": "create"})
    assert result["counts"]["created"] == 1
    [created] = _by_name(db, tag_c)
    assert created.data["allowDuplicate"] is True and "importDuplicateOverride" in created.data
    assert db.get(Candidate, first["id"]).phone == phone and db.get(Candidate, second["id"]).email == email   # fiches intactes


def test_explicit_update_shows_changes_and_never_erases_with_an_empty_cell(client, auth_headers, db):
    tag = _tag()
    phone = _phone()
    existing = client.post("/api/drh/candidates", headers=auth_headers, json={
        "first_name": "Karim", "last_name": f"IMPORT{tag}", "phone": phone, "email": f"old{tag.lower()}@example.com",
        "desired_position": "Chauffeur", "status": "nouvelle",
        "data": {"wilaya": "Oran", "notes": "À conserver", "avisDecision": "Instance", "sexe": "M"}}).json()["data"]
    row = _row(tag, **{"Téléphone": phone, "E-mail": None, "Wilaya": "Alger", "Commune": "Alger Centre", "Sexe": None, "Date de naissance": "15/03/1990"})
    upload = _upload(client, auth_headers, _xlsx([row]))
    line = _preview(client, auth_headers, upload).json()["rows"][0]
    assert line["actions"] == ["skip", "create", "update"] and line["default_action"] == "skip" and line["update_target"] == existing["id"]
    assert {change["field"]: (change["old"], change["new"]) for change in line["changes"]} == {
        "birth_date": ("", "15/03/1990"), "wilaya": ("Oran", "Alger"), "commune": ("", "Alger Centre")}
    result = _confirm(client, auth_headers, upload.json()["session_id"], {"2": "update"}).json()
    assert result["counts"]["updated"] == 1 and result["counts"]["created"] == 0
    assert result["rows"][0]["candidate_id"] == existing["id"]
    db.expire_all()
    updated = db.get(Candidate, existing["id"])
    assert updated.email == f"old{tag.lower()}@example.com"                         # cellule vide : valeur conservée
    assert updated.data["sexe"] == "M" and updated.data["notes"] == "À conserver" and updated.data["avisDecision"] == "Instance"
    assert updated.data["wilaya"] == "Alger" and updated.data["commune"] == "Alger Centre" and updated.data["dateNaissance"] == "1990-03-15"
    assert updated.status == "nouvelle" and updated.society is None and updated.data["auditTrail"][-1]["action"] == "modification"
    assert len(_by_name(db, tag)) == 1


def test_update_is_refused_without_the_update_right(client, auth_headers, restricted_headers, db):
    tag = _tag()
    phone = _phone()
    existing = client.post("/api/drh/candidates", headers=auth_headers, json={
        "first_name": "Karim", "last_name": f"IMPORT{tag}", "phone": phone, "status": "nouvelle", "data": {"wilaya": "Oran"}}).json()["data"]
    creator = _recruiter(client, db, actions=("read", "create"))
    config = client.get(f"{BASE}/config", headers=creator).json()
    assert config["can_update"] is False and config["limits"]["max_rows"] == importer.MAX_ROWS
    upload = _upload(client, creator, _xlsx([_row(tag, **{"Téléphone": phone})]))
    line = _preview(client, creator, upload).json()["rows"][0]
    assert line["actions"] == ["skip", "create"] and line["changes"] == [] and line["update_target"] is None
    result = _confirm(client, creator, upload.json()["session_id"], {"2": "update"}).json()     # requête forgée
    assert result["counts"]["updated"] == 0 and result["counts"]["failed"] == 1 and "non autorisée" in result["rows"][0]["reason"]
    db.expire_all()
    assert db.get(Candidate, existing["id"]).data["wilaya"] == "Oran"
    # Lecture seule : aucun accès à l'import.
    reader = _recruiter(client, db, actions=("read",))
    assert _upload(client, reader, _xlsx([_row(_tag())])).status_code == 403
    assert client.get(f"{BASE}/template", headers=reader).status_code == 403
    # Compte sans module Recrutement.
    assert _upload(client, restricted_headers, _xlsx([_row(_tag())])).status_code == 403


# ── Périmètre société ────────────────────────────────────────────────────────────────────
def test_matches_outside_the_society_scope_are_neither_exposed_nor_modified(client, auth_headers, db):
    tag, tag_emp, tag_pool = _tag(), _tag(), _tag()
    phone, emp_phone, pool_phone = _phone(), _phone(), _phone()
    # Dossier transféré à la DRH d'une autre société : sorti du vivier Recrutement.
    foreign = Candidate(first_name="SECRET", last_name=f"IMPORT{tag}", phone=phone, society=SOLUTION, status="a_contractualiser", data={})
    employee = Employee(code=f"E{tag_emp}", first_name="AGENT", last_name=f"IMPORT{tag_emp}", phone=emp_phone, society=SOLUTION, status="actif")
    ventilated = Candidate(first_name="VIVIER", last_name=f"IMPORT{tag_pool}", phone=pool_phone, society=SOLUTION, status="nouvelle", data={"wilaya": "Oran"})
    db.add_all([foreign, employee, ventilated])
    db.commit()
    recruiter = _recruiter(client, db, societies=(SCOPE_SECURITE,))
    rows = [_row(_tag(), **{"Téléphone": phone}), _row(_tag(), **{"Téléphone": emp_phone}), _row(_tag(), **{"Téléphone": pool_phone, "Wilaya": "Alger"})]
    upload = _upload(client, recruiter, _xlsx(rows))
    preview = _preview(client, recruiter, upload)
    lines = preview.json()["rows"]
    hidden = {"type": "candidat", "reasons": ["telephone"], "visible": False, "updatable": False}
    assert lines[0]["matches"] == [hidden] and lines[1]["matches"] == [{**hidden, "type": "salarie"}]
    assert "SECRET" not in preview.text and f"IMPORT{tag}" not in preview.text and f"IMPORT{tag_emp}" not in preview.text
    assert str(foreign.id) not in {str(match.get("id")) for line in lines[:2] for match in line["matches"]}
    assert all(line["actions"] == ["skip", "create"] and line["default_action"] == "skip" for line in lines)
    # Vivier Groupe : le dossier ventilé ailleurs reste visible (règle existante) mais non modifiable.
    assert lines[2]["matches"][0]["visible"] is True and lines[2]["matches"][0]["updatable"] is False and lines[2]["update_target"] is None
    result = _confirm(client, recruiter, upload.json()["session_id"], {"2": "update", "4": "update"}).json()
    assert result["counts"]["updated"] == 0 and result["counts"]["created"] == 0
    assert [row["outcome"] for row in result["rows"]] == ["failed", "skipped", "failed"]
    assert "hors de votre périmètre" in result["rows"][1]["reason"] and "n°" not in result["rows"][1]["reason"]
    db.expire_all()
    assert db.get(Candidate, ventilated.id).data == {"wilaya": "Oran"} and db.get(Candidate, foreign.id).phone == phone
    # Le même fichier vu par un compte global : la fiche est détaillée.
    admin_lines = _preview(client, auth_headers, _upload(client, auth_headers, _xlsx(rows))).json()["rows"]
    assert admin_lines[0]["matches"][0]["id"] == foreign.id and admin_lines[0]["matches"][0]["state"] == "Transmis à la DRH"
    assert admin_lines[1]["matches"][0]["state"] == "Salarié"


def test_import_session_belongs_to_its_user_and_scope(client, auth_headers, db):
    owner = _recruiter(client, db)
    other = _recruiter(client, db)
    tag = _tag()
    upload = _upload(client, owner, _xlsx([_row(tag)]))
    session_id = upload.json()["session_id"]
    assert _preview(client, owner, upload).status_code == 200
    for call in (lambda: client.post(f"{BASE}/{session_id}/preview", headers=other, json={"sheet": 0, "mapping": {}}),
                 lambda: _confirm(client, other, session_id), lambda: _confirm(client, auth_headers, session_id),
                 lambda: client.get(f"{BASE}/{session_id}/report", headers=other),
                 lambda: client.post(f"{BASE}/{session_id}/cancel", headers=other)):
        assert call().status_code == 404
    assert _by_name(db, tag) == []
    # Périmètre modifié entre l'analyse et la confirmation : la session n'est plus utilisable.
    session = db.execute(select(CandidateImportSession).where(CandidateImportSession.public_id == session_id)).scalar_one()
    user = db.get(User, session.user_id)
    user.authorized_societies = [SCOPE_SOLUTION]
    db.commit()
    changed = _confirm(client, owner, session_id)
    assert changed.status_code == 409 and "périmètre" in changed.json()["detail"]
    assert _by_name(db, tag) == []


# ── Fiabilité : annulation, double confirmation, concurrence ─────────────────────────────
def test_cancel_or_abandon_before_confirmation_creates_nothing(client, auth_headers, db):
    tag = _tag()
    before = db.query(Candidate).count()
    upload = _upload(client, auth_headers, _xlsx([_row(tag), _row(_tag())]))
    session_id = upload.json()["session_id"]
    assert _preview(client, auth_headers, upload).status_code == 200
    assert db.query(Candidate).count() == before                                    # analyse et vérification : rien créé
    assert client.post(f"{BASE}/{session_id}/cancel", headers=auth_headers).json() == {"status": "cancelled"}
    late = _confirm(client, auth_headers, session_id)
    assert late.status_code == 410
    db.expire_all()
    session = db.execute(select(CandidateImportSession).where(CandidateImportSession.public_id == session_id)).scalar_one()
    assert session.status == "cancelled" and session.workbook is None and session.plan is None   # contenu effacé
    assert db.query(Candidate).count() == before and _by_name(db, tag) == []
    # Confirmer sans vérification préalable est refusé.
    fresh = _upload(client, auth_headers, _xlsx([_row(tag)]))
    assert _confirm(client, auth_headers, fresh.json()["session_id"]).status_code == 409
    assert _by_name(db, tag) == []


def test_double_confirmation_and_retry_never_create_twice(client, auth_headers, db):
    tag = _tag()
    content = _xlsx([_row(tag)])
    upload = _upload(client, auth_headers, content)
    session_id = upload.json()["session_id"]
    _preview(client, auth_headers, upload)
    first = _confirm(client, auth_headers, session_id).json()
    second = _confirm(client, auth_headers, session_id)                             # double clic / nouvelle tentative
    assert first["counts"]["created"] == 1 and first["already_processed"] is False
    assert second.status_code == 200 and second.json()["already_processed"] is True
    assert second.json()["counts"] == first["counts"] and second.json()["rows"] == first["rows"]
    assert len(_by_name(db, tag)) == 1
    db.expire_all()
    session = db.execute(select(CandidateImportSession).where(CandidateImportSession.public_id == session_id)).scalar_one()
    assert session.status == "completed" and session.workbook is None               # contenu du fichier effacé
    assert "IMPORT" + tag not in str(session.result)                                # résultat sans contenu personnel
    # Nouvel envoi du même fichier : la fiche existe déjà, elle est ignorée.
    _, again = _run(client, auth_headers, content)
    assert again["counts"]["created"] == 0 and again["counts"]["skipped"] == 1
    assert len(_by_name(db, tag)) == 1


def test_concurrent_confirmations_of_one_session_create_once(live_client, live_headers):
    from tests.conftest import TestSessionLocal

    tag = _tag()
    with TestSessionLocal() as session:
        if not session.query(Position).filter(Position.name == "Chauffeur").first():
            session.add(Position(name="Chauffeur", society=None))
            session.commit()
    upload = _upload(live_client, live_headers, _xlsx([_row(tag), _row(_tag()), _row(_tag())]))
    session_id = upload.json()["session_id"]
    assert _preview(live_client, live_headers, upload).status_code == 200
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(lambda _: _confirm(live_client, live_headers, session_id), range(4)))
    assert {response.status_code for response in responses} <= {200, 409}
    fresh = [response.json() for response in responses if response.status_code == 200 and not response.json()["already_processed"]]
    assert len(fresh) == 1 and fresh[0]["counts"]["created"] == 3
    with TestSessionLocal() as session:
        assert session.query(Candidate).filter(Candidate.last_name == f"IMPORT{tag}").count() == 1


def test_duplicates_are_rechecked_at_confirmation(client, auth_headers, db):
    tag, other = _tag(), _tag()
    phone = _phone()
    upload = _upload(client, auth_headers, _xlsx([_row(tag, **{"Téléphone": phone}), _row(other)]))
    assert _preview(client, auth_headers, upload).json()["counts"]["valid"] == 2
    # Création concurrente entre la vérification et la confirmation.
    created = client.post("/api/drh/candidates", headers=auth_headers, json={"first_name": "Autre", "last_name": f"CONC{tag}", "phone": phone, "status": "nouvelle"})
    assert created.status_code == 200
    result = _confirm(client, auth_headers, upload.json()["session_id"]).json()
    assert result["counts"]["created"] == 1 and result["counts"]["skipped"] == 1
    assert "création concurrente" in result["rows"][0]["reason"]
    assert _by_name(db, tag) == [] and len(_by_name(db, other)) == 1


def test_expired_sessions_are_purged(client, auth_headers, db):
    tag = _tag()
    upload = _upload(client, auth_headers, _xlsx([_row(tag)]))
    session_id = upload.json()["session_id"]
    _preview(client, auth_headers, upload)
    session = db.execute(select(CandidateImportSession).where(CandidateImportSession.public_id == session_id)).scalar_one()
    session.expires_at = datetime.utcnow() - timedelta(minutes=1)
    db.commit()
    expired = _confirm(client, auth_headers, session_id)
    assert expired.status_code == 410 and "expirée" in expired.json()["detail"]
    db.expire_all()
    assert session.status == "expired" and session.workbook is None and _by_name(db, tag) == []
    # Une nouvelle analyse abandonne et vide les analyses en attente du même compte.
    first = _upload(client, auth_headers, _xlsx([_row(_tag())])).json()["session_id"]
    _upload(client, auth_headers, _xlsx([_row(_tag())]))
    db.expire_all()
    previous = db.execute(select(CandidateImportSession).where(CandidateImportSession.public_id == first)).scalar_one()
    assert previous.status == "cancelled" and previous.workbook is None


# ── Fichiers invalides et limites ────────────────────────────────────────────────────────
def test_invalid_files_and_limits_are_refused_with_a_clear_reason(client, auth_headers, db):
    before = db.query(Candidate).count()
    encrypted = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64 + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 64
    cases = [
        ("candidats.csv", b"Nom;Prenom\nA;B\n", "Format non pris en charge"),
        ("candidats.xlsm", _xlsx([_row(_tag())]), "Format non pris en charge"),      # classeur à macros refusé
        ("candidats.xlsx", b"", "vide"),
        ("candidats.xlsx", b"PK\x03\x04" + b"corrompu" * 40, "corrompu"),
        ("candidats.xlsx", b"ceci n'est pas un classeur", "Format incorrect"),
        ("candidats.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x01" * 600, "corrompu"),
        ("candidats.xlsx", encrypted, "protégé par mot de passe"),
        ("candidats.xls", encrypted, "protégé par mot de passe"),
        ("candidats.xlsx", b"PK\x03\x04" + b"0" * (importer.MAX_FILE_BYTES + 10), "trop volumineux"),
    ]
    for name, content, expected in cases:
        response = _upload(client, auth_headers, content, name)
        assert response.status_code == 422 and expected in response.json()["detail"], (name, response.text)
    # Trop de lignes : la feuille est signalée, jamais tronquée en silence.
    many = _upload(client, auth_headers, _xlsx([[f"NOM{i}", "Test"] for i in range(importer.MAX_ROWS + 1)], ["Nom", "Prénom"]))
    sheet = many.json()["sheets"][0]
    assert sheet["importable"] is False and str(importer.MAX_ROWS) in sheet["reason"] and many.json()["default_sheet"] is None
    assert _preview(client, auth_headers, many, sheet=0, mapping={"0": "last_name", "1": "first_name"}).status_code == 422
    assert db.query(Candidate).count() == before
    refused = db.execute(select(AuditEvent).where(AuditEvent.action == "recruitment.candidates.import.analyze", AuditEvent.result == "refused")).scalars().all()
    assert len(refused) >= len(cases)


# ── Journal, modèle, rapport ─────────────────────────────────────────────────────────────
def test_audit_records_author_scope_file_and_counters_without_personal_content(client, auth_headers, db):
    tag = _tag()
    phone = _phone()
    upload = _upload(client, auth_headers, _xlsx([_row(tag, **{"Téléphone": phone}), _row(_tag(), **{"E-mail": "invalide"})]), "vivier-avril.xlsx")
    session_id = upload.json()["session_id"]
    _preview(client, auth_headers, upload)
    _confirm(client, auth_headers, session_id)
    db.expire_all()
    [event] = db.execute(select(AuditEvent).where(AuditEvent.action == "recruitment.candidates.import", AuditEvent.resource_id == session_id)).scalars().all()
    assert event.username == "testadmin" and event.result == "success" and event.society == "Toutes sociétés" and event.created_at
    assert '"file_name": "vivier-avril.xlsx"' in event.new_state and '"created": 1' in event.new_state and '"errors": 1' in event.new_state
    for personal in (f"IMPORT{tag}", phone, "example.com"):
        assert personal not in event.new_state


def test_template_matches_the_importable_fields(client, auth_headers, db):
    response = client.get(f"{BASE}/template", headers=auth_headers)
    assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Candidats", "Listes", "Exemple", "Instructions"]
    sheet = book["Candidats"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == [item.label + (" *" if item.required else "") for item in FIELDS]
    assert headers[:2] == ["Nom *", "Prénom *"] and sheet.max_row >= 2 and sheet["A2"].value is None   # aucune ligne d'exemple à importer
    assert [field_for_header(header) for header in headers] == [item.key for item in FIELDS]           # relu sans correction
    phone_column = next(cell.column_letter for cell in sheet[1] if cell.value == "Téléphone")
    assert sheet[f"{phone_column}2"].number_format == "@" and sheet[f"{phone_column}{importer.MAX_ROWS + 1}"].number_format == "@"
    assert sheet["A1"].fill.fgColor.rgb != sheet["C1"].fill.fgColor.rgb                               # obligatoires distingués
    validated = {str(validation.sqref) for validation in sheet.data_validations.dataValidation}
    sex_column = next(cell.column_letter for cell in sheet[1] if cell.value == "Sexe")
    assert f"{sex_column}2:{sex_column}{importer.MAX_ROWS + 1}" in validated
    assert "Chauffeur" in [cell.value for column in book["Listes"].iter_cols() for cell in column]
    assert "ne pas importer" in book["Exemple"]["A1"].value and book["Exemple"]["A4"].value == "BENALI"
    assert any("Texte" in str(row[0].value) for row in book["Instructions"].iter_rows())
    # Le modèle téléversé tel quel : seule la feuille « Candidats » est proposée, et elle est vide.
    upload = _upload(client, auth_headers, response.content, "modele-import-candidats.xlsx").json()
    assert [(sheet["name"], sheet["importable"]) for sheet in upload["sheets"]] == [
        ("Candidats", False), ("Listes", False), ("Exemple", False), ("Instructions", False)]
    # Rempli puis réimporté : toutes les colonnes sont reconnues.
    tag = _tag()
    values = {"Nom *": f"IMPORT{tag}", "Prénom *": "Lina", "Téléphone": _phone(), "Sexe": "F", "Wilaya": "Alger", "Commune": "Alger Centre",
              "Poste souhaité": "Chauffeur", "Situation familiale": "Marié(e)", "Groupe sanguin": "O+", "Disponibilité": "Immédiatement",
              "Langues parlées": "Arabe, Français", "Observations": "RAS", "Source de candidature": "ANEM", "Taille (cm)": 170,
              "NIN": "109900123456789012", "Salaire demandé (DA/mois)": "45 000,00 DA", "Expérience professionnelle": "3 ans"}
    for column, header in enumerate(headers, start=1):
        if header in values:
            sheet.cell(row=2, column=column, value=values[header])
    buffer = io.BytesIO()
    book.save(buffer)
    preview, result = _run(client, auth_headers, buffer.getvalue())
    assert preview["ignored_columns"] == [] and result["counts"]["created"] == 1
    [candidate] = _by_name(db, tag)
    assert candidate.expected_salary == 45000 and candidate.data["langues"] == ["Arabe", "Français"] and candidate.data["taille"] == "170"
    assert candidate.data["nin"] == "109900123456789012" and candidate.data["groupeSanguin"] == "O+" and candidate.data["experienceTexte"] == "3 ans"
    assert candidate.data["situation"] == "Marié(e)" and candidate.data["source"] == "ANEM" and candidate.data["notes"] == "RAS"


def test_error_report_lists_rows_and_reasons_with_formulas_neutralised(client, auth_headers, db):
    tag = _tag()
    rows = [_row(tag), _row(_tag(), **{"Poste souhaité": "=HYPERLINK(\"http://x\")"}), _row(_tag(), **{"Sexe": "@cmd"}), _row(_tag(), **{"E-mail": "+bad"})]
    upload = _upload(client, auth_headers, _xlsx(rows), "=cmd.xlsx")
    session_id = upload.json()["session_id"]
    _preview(client, auth_headers, upload)
    before = load_workbook(io.BytesIO(client.get(f"{BASE}/{session_id}/report", headers=auth_headers).content)).active   # avant confirmation
    assert any(row[2].value == "Erreur de données" for row in before.iter_rows())
    result = _confirm(client, auth_headers, session_id).json()
    assert result["counts"]["created"] == 1 and result["counts"]["errors"] == 3
    response = client.get(f"{BASE}/{session_id}/report", headers=auth_headers)
    assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    cells = [cell for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert all(cell.data_type != "f" for cell in cells)                              # aucune formule dans l'export
    assert all(not str(cell.value).startswith(("=", "+", "-", "@")) for cell in cells if isinstance(cell.value, str))
    table = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert ["Fichier", "'=cmd.xlsx"] in [row[:2] for row in table]
    assert ["Candidats créés", 1] in [row[:2] for row in table] and ["Lignes en erreur", 3] in [row[:2] for row in table]
    details = {row[1]: row for row in table if row[2] in ("Créé", "Erreur de données")}
    assert details[2][2] == "Créé" and details[2][5] == _by_name(db, tag)[0].id
    assert details[3][3] == "Poste souhaité" and "formule" in details[3][4]          # formule du fichier jamais évaluée
    assert details[4][3] == "Sexe" and details[5][3] == "E-mail"
    assert importer.neutralize("=1+1") == "'=1+1" and importer.neutralize("-2") == "'-2" and importer.neutralize("Alger") == "Alger"
    assert importer.neutralize(12) == 12


def test_reference_lists_match_the_candidate_form():
    """Le référentiel Python reprend exactement les listes de la fiche candidat (recrute.html)."""
    import re
    from pathlib import Path

    html = (Path(__file__).resolve().parents[1] / "app" / "static" / "recrute.html").read_text(encoding="utf-8")
    block = re.search(r"const RECRUTE_WILAYAS=\[(.*?)\];", html, re.S).group(1)
    assert tuple(re.findall(r'\["(\d{2})","([^"]+)"\]', block)) == WILAYAS

    def options(field_id):
        select_html = re.search(rf'<select id="{field_id}"[^>]*>(.*?)</select>', html, re.S).group(1)
        found = re.findall(r'<option(?: value="([^"]*)")?>([^<]*)</option>', select_html)
        return [value or text for value, text in found if value or not text.startswith("—")]

    by_key = {item.key: item for item in FIELDS}
    for key, field_id in (("family_status", "fFamilyStatus"), ("blood_group", "fBlood"), ("availability", "fAvailability"),
                          ("military_service", "fMilitary"), ("shirt_size", "fShirtSize"), ("sex", "fSex")):
        expected = [value for value in options(field_id) if value not in ("—", "— Choisir —")]
        assert by_key[key].choice_values() == expected, key
    for label in ("Nom", "Prénom", "Date de naissance", "Lieu de naissance", "Sexe", "Situation familiale", "Téléphone", "Adresse",
                  "Commune", "Wilaya", "Poste souhaité", "Disponibilité", "Source de candidature", "Service militaire", "NIN"):
        assert f">{label}</label>" in html and any(norm_key(item.label) == norm_key(label) for item in FIELDS), label
