"""Tests du vrai parcours upload → preview → confirmation (aucun mock Excel)."""
from datetime import date
from io import BytesIO
import uuid

import pytest
from docx import Document
from openpyxl import Workbook
from sqlalchemy import select

from app.modules.auth.models import AuditEvent
from app.modules.commercial.models import Client
from app.modules.drh import contract_import, service
from app.modules.drh.models import Candidate, Contract, ContractTemplate, Employee, GeneratedContract
from app.modules.drh.candidate_import_models import CandidateImportSession

BASE = '/api/drh/contracts/import'


@pytest.fixture
def setup_import(db):
    society = 'IMPORT-' + uuid.uuid4().hex[:8]
    dc = Client(name='CLIENT EXCEL', society=society, status='actif')
    doc = Document(); doc.add_paragraph('{{NOM}} {{PRENOM}} {{DATE_DEBUT}}')
    content = BytesIO(); doc.save(content)
    template = ContractTemplate(code=uuid.uuid4().hex, title='Modèle test', contract_type='CDD', active=1, file_name='modele.docx', docx_content=content.getvalue())
    candidates = [Candidate(first_name=f'Prenom{i}', last_name='EXCEL', society=society, status='a_contractualiser', data={'avisDecision':'Favorable'}, phone=f'055012300{i}') for i in range(3)]
    db.add_all([dc, template, *candidates]); db.commit()
    return society, candidates, {'position':'MAGASINIER','contract_type':'CDD','duration':'12m','client_site':f'{dc.id}:','template_id':str(template.id),'salary':'40000'}


def upload(client, headers, society, rows, titles=None):
    book=Workbook(); sheet=book.active
    sheet.append(titles or ['Référence candidat','Date début'])
    for row in rows: sheet.append(row)
    content=BytesIO(); book.save(content)
    result=client.post(BASE, params={'society':society}, headers=headers, files={'file':('contrats.xlsx', content.getvalue())})
    assert result.status_code==200, result.text
    return result.json()['session_id']


def preview(client, headers, sid, common, **config):
    result=client.post(f'{BASE}/{sid}/preview', headers=headers, json={'sheet':0, 'mapping':{'0':'reference','1':'start_date'},'common':common, **config})
    assert result.status_code==200, result.text
    return result.json()


def confirm(client, headers, sid):
    response=client.post(f'{BASE}/{sid}/confirm',headers=headers)
    assert response.status_code==200, response.text
    return response.json()


def test_dates_only_distinct_contracts_and_audit(client, auth_headers, db, setup_import):
    society, candidates, common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15'],[candidates[1].id,date(2026,11,20)]])
    result=preview(client,auth_headers,sid,common)
    assert result['counts']=={'PRÊT':2}, result
    assert [r['values']['end_date'] for r in result['rows']]==['2027-10-15','2027-11-20']
    assert not db.scalars(select(Employee).where(Employee.society==society)).all()
    assert not db.scalars(select(Contract).join(Employee).where(Employee.society==society)).all()
    result=confirm(client,auth_headers,sid)
    assert result['counts']=={'CRÉÉ':2}
    assert len({r['reference'] for r in result['rows']})==2
    contracts=db.scalars(select(Contract).join(Employee).where(Employee.society==society)).all()
    assert len(contracts)==2
    assert {c.start_date for c in contracts}=={date(2026,10,15),date(2026,11,20)}
    assert all(c.dc_client_id and c.template_code for c in contracts)
    assert len(db.scalars(select(AuditEvent).where(AuditEvent.action=='drh.contract.excel.created',AuditEvent.society==society)).all())==2
    assert confirm(client,auth_headers,sid)==result  # confirmation idempotente
    db.expire_all()
    assert db.get(CandidateImportSession, db.scalar(select(CandidateImportSession.id).where(CandidateImportSession.public_id==sid))).workbook is None


def test_full_import_and_explicit_precedence(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    titles=['Référence candidat','Date début','Poste / fonction','Salaire','Banque','RIB / IBAN','Date fin manuelle']
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-01-31','CARISTE',50000,'BEA','001234','2026-12-25']],titles)
    mapping={'0':'reference','1':'start_date','2':'position','3':'salary','4':'bank','5':'iban','6':'end_date'}
    result=preview(client,auth_headers,sid,common,mode='full',precedence='excel',mapping=mapping)
    v=result['rows'][0]['values']
    assert v['position']=='CARISTE' and v['salary']==50000 and v['end_date']=='2026-12-25'
    result=preview(client,auth_headers,sid,common,mode='full',precedence='common',mapping=mapping)
    assert result['rows'][0]['values']['position']=='MAGASINIER'
    assert confirm(client,auth_headers,sid)['counts']=={'CRÉÉ':1}
    employee=db.scalar(select(Employee).where(Employee.extra['sourceCandidateId'].as_integer()==candidates[0].id))
    assert employee.extra['banque']=='BEA' and employee.extra['rib']=='001234'


@pytest.mark.parametrize('change,status', [
    ({'position':''},'DONNÉE MANQUANTE'), ({'client_site':'inconnu'},'À CORRIGER'),
    ({'template_id':'99999999'},'À CORRIGER'), ({'duration':'invalide'},'À CORRIGER'),
    ({'salary':'-1'},'À CORRIGER'), ({'contract_type':'INCONNU'},'À CORRIGER'),
])
def test_invalid_values(client,auth_headers,setup_import,change,status):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    result=preview(client,auth_headers,sid,{**common,**change})
    assert result['rows'][0]['status']==status
    assert confirm(client,auth_headers,sid)['rows'][0]['status']==status


def test_missing_duplicate_already_contracted(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    candidates[2].status='embauche';db.commit()
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15'],[candidates[0].id,'2026-10-16'],['missing','2026-10-15'],[candidates[2].id,'2026-10-15']])
    result=preview(client,auth_headers,sid,common)
    assert [r['status'] for r in result['rows']]==['DOUBLON','DOUBLON','CANDIDAT INTROUVABLE','DÉJÀ CONTRACTUALISÉ']


def test_ambiguous_phone_is_not_automatic(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    candidates[1].phone=candidates[0].phone;db.commit()
    sid=upload(client,auth_headers,society,[[candidates[0].phone,'2026-10-15']],['Téléphone','Date début'])
    result=preview(client,auth_headers,sid,common,mapping={'0':'phone','1':'start_date'})
    assert result['rows'][0]['status']=='À VÉRIFIER'


def test_partial_failure_rolls_back_candidate_and_keeps_others(client,auth_headers,db,setup_import,monkeypatch):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[c.id,'2026-10-15'] for c in candidates])
    preview(client,auth_headers,sid,common)
    original=service.generate_contract
    def fail_one(db,request,*args,**kwargs):
        employee=db.get(Employee,request.employee_id)
        if (employee.extra or {}).get('sourceCandidateId')==candidates[1].id: raise RuntimeError('failure rendering')
        return original(db,request,*args,**kwargs)
    monkeypatch.setattr(service,'generate_contract',fail_one)
    result=confirm(client,auth_headers,sid)
    assert result['counts']=={'CRÉÉ':2,'ÉCHEC':1}
    db.expire_all()
    assert db.get(Candidate,candidates[1].id).status=='a_contractualiser'
    assert len(db.scalars(select(Employee).where(Employee.society==society)).all())==2
    assert len(db.scalars(select(GeneratedContract).join(Employee).where(Employee.society==society)).all())==2


def test_revalidation_before_confirmation(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    preview(client,auth_headers,sid,common)
    candidates[0].status='embauche';db.commit()
    assert confirm(client,auth_headers,sid)['counts']=={'DÉJÀ CONTRACTUALISÉ':1}


def test_confirm_requires_preview(client,auth_headers,setup_import):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    assert client.post(f'{BASE}/{sid}/confirm',headers=auth_headers).status_code==409


def test_scope_and_permissions(client,auth_headers,restricted_headers,setup_import):
    society,candidates,common=setup_import
    assert client.get(BASE+'/config',headers=restricted_headers).status_code==403
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    assert client.post(f'{BASE}/{sid}/preview',headers=restricted_headers,json={'sheet':0,'mapping':{}}).status_code==403


def test_candidate_outside_active_society_not_found(client,auth_headers,setup_import):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,'OTHER-'+society,[[candidates[0].id,'2026-10-15']])
    result=preview(client,auth_headers,sid,common)
    assert result['rows'][0]['status']=='CANDIDAT INTROUVABLE'


def test_report_formula_injection_and_session_isolation(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    candidates[0].last_name='=DANGER';db.commit()
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    preview(client,auth_headers,sid,common)
    report=client.get(f'{BASE}/{sid}/report',headers=auth_headers)
    assert report.status_code==200 and "'=DANGER" in report.text
    response=client.post(f'/api/drh/candidates/import/{sid}/preview',headers=auth_headers,json={'sheet':0,'mapping':{}})
    assert response.status_code==404


def test_duration_matches_existing_month_overflow():
    assert contract_import.duration_end(date(2026,1,31),'1m')==date(2026,3,3)
    assert contract_import.duration_end(date(2026,1,31),'30d')==date(2026,3,2)

@pytest.mark.parametrize('start,end', [('31/02/2026',None), ('',None), ('2026-10-15','2026-10-01')])
def test_bad_or_missing_dates(client,auth_headers,setup_import,start,end):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,start,end]],['Référence candidat','Date début','Date fin manuelle'])
    result=preview(client,auth_headers,sid,common,mapping={'0':'reference','1':'start_date','2':'end_date'})
    assert result['rows'][0]['status']=='À CORRIGER'


def test_contract_and_candidate_sessions_coexist(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    from app.modules.auth.models import User
    from app.modules.drh import candidate_import
    book=Workbook();book.active.append(['Nom','Prénom']);book.active.append(['TEST','Candidat'])
    content=BytesIO();book.save(content)
    user=db.scalar(select(User).where(User.username=='testadmin'))
    candidate_import.create_session(db,user,file_name='candidats.xlsx',content=content.getvalue())
    db.expire_all()
    assert db.scalar(select(CandidateImportSession.status).where(CandidateImportSession.public_id==sid))=='uploaded'
    assert preview(client,auth_headers,sid,common)['counts']=={'PRÊT':1}


def test_readonly_drh_and_other_user_session_denied(client,auth_headers,db,setup_import):
    from app.core.security import hash_password
    from app.modules.auth.models import User
    society,candidates,common=setup_import
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    username='import-'+uuid.uuid4().hex[:8]
    user=User(username=username,email=username+'@test.com',full_name='DRH',role='rh',access_level='H3',authorized_modules=['drh'],authorized_societies=[society],authorized_actions=['read'],is_active=True,password_hash=hash_password('testpass123'))
    db.add(user);db.commit()
    token=client.post('/api/auth/login',json={'username':username,'password':'testpass123'}).json()['access_token']
    headers={'Authorization':'Bearer '+token}
    assert client.get(BASE+'/config',headers=headers).status_code==403
    user.authorized_actions=['read','create'];db.commit()
    assert client.post(f'{BASE}/{sid}/preview',headers=headers,json={'sheet':0,'mapping':{}}).status_code==404
    book=Workbook();book.active.append(['Référence candidat','Date début']);book.active.append([candidates[0].id,'2026-10-15']);content=BytesIO();book.save(content)
    response=client.post(BASE,params={'society':'SOCIETE INTERDITE'},headers=headers,files={'file':('contrats.xlsx',content.getvalue())})
    assert response.status_code==403


def test_unready_row_never_promoted_silently(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    candidates[0].status='reserve';db.commit()
    sid=upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']])
    assert preview(client,auth_headers,sid,common)['counts']=={'À CORRIGER':1}
    candidates[0].status='a_contractualiser';db.commit()
    assert confirm(client,auth_headers,sid)['counts']=={'NON CRÉÉ':1}
    assert not db.scalar(select(Employee.id).where(Employee.society==society))


def test_full_excel_without_common_values(client,auth_headers,db,setup_import):
    society,candidates,common=setup_import
    keys=['reference','start_date','position','client_site','contract_type','duration','trial','template_id','salary','payment','bank','iban']
    values=[candidates[0].id,'2026-10-15','CARISTE',common['client_site'],'CDD','6m','15d',common['template_id'],'55000','Virement','BEA','0012345678']
    sid=upload(client,auth_headers,society,[values],[contract_import.FIELDS[k] for k in keys])
    result=preview(client,auth_headers,sid,{},mode='full',precedence='excel',mapping={str(i):key for i,key in enumerate(keys)})
    assert result['counts']=={'PRÊT':1}
    assert result['rows'][0]['values']['end_date']=='2027-04-15'
    assert confirm(client,auth_headers,sid)['counts']=={'CRÉÉ':1}
    employee=db.scalar(select(Employee).where(Employee.extra['sourceCandidateId'].as_integer()==candidates[0].id))
    assert employee.trial_end_date==date(2026,10,30)
    assert employee.extra['modePaiement']=='Virement' and employee.extra['rib']=='0012345678'
    db.expire_all()
    assert db.get(Candidate,candidates[0].id).data['convertedContractId']


def test_two_imports_cannot_create_two_contracts_for_one_candidate(client,auth_headers,db,setup_import):
    from concurrent.futures import ThreadPoolExecutor
    society,candidates,common=setup_import
    sessions=[upload(client,auth_headers,society,[[candidates[0].id,'2026-10-15']]) for _ in range(2)]
    for sid in sessions:
        assert preview(client,auth_headers,sid,common)['counts']=={'PRÊT':1}
    from tests.conftest import TestSessionLocal
    from app.main import app
    from app.db.session import get_db
    previous = app.dependency_overrides[get_db]
    def separate_session():
        with TestSessionLocal() as session:
            yield session
    app.dependency_overrides[get_db] = separate_session
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses=list(pool.map(lambda sid: client.post(f'{BASE}/{sid}/confirm',headers=auth_headers),sessions))
    finally:
        app.dependency_overrides[get_db] = previous
    assert all(r.status_code==200 for r in responses), [r.text for r in responses]
    assert sum(r.json()['counts'].get('CRÉÉ',0) for r in responses)==1
    db.expire_all()
    employees=db.scalars(select(Employee).where(Employee.society==society)).all()
    assert len(employees)==1
    assert len(db.scalars(select(Contract).where(Contract.employee_id==employees[0].id)).all())==1
