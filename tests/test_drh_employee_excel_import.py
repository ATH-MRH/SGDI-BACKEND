"""Import Excel des employés en masse : vrai parcours upload → aperçu → corrections → confirmation.

Aucun mock Excel : classeurs openpyxl réels, base de test, API réelle. Les identités sont
uniques par test (la base de test est partagée par toute la session pytest)."""
from datetime import date
from io import BytesIO
import random
import string
import uuid

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.modules.auth.models import AuditEvent, User
from app.modules.commercial.models import Client
from app.modules.drh import candidate_import, employee_import
from app.modules.drh.candidate_import_models import CandidateImportSession
from app.modules.drh.models import Contract, Employee
from app.modules.irongs.models import Position
from app.modules.ops.models import Assignment, Site

BASE = "/api/drh/employee-import"
TITLES = ["Matricule", "Nom", "Prénom", "Téléphone", "Poste", "Date de recrutement"]


def letters(size=6):
    return "".join(random.choices(string.ascii_uppercase, k=size))


def phone():
    return "055" + "".join(random.choices(string.digits, k=7))


@pytest.fixture
def system_headers(client):
    response = client.post("/api/auth/admin-system-login", json={"username": "testadmin", "password": "test-admin-password"})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture
def ref(db):
    """Société connue (par un employé existant), poste, client et deux sites actifs."""
    tag = letters()
    society = f"SOC EMP {tag}"
    seed = Employee(code=f"SEED{tag}", first_name="SEED", last_name=f"SEED{tag}", society=society, phone=phone(), status="actif")
    position = Position(name=f"AGENT {tag}", society=None, active=True)
    client_row = Client(name=f"CLIENT {tag}", society=society, status="actif")
    db.add_all([seed, position, client_row])
    db.flush()
    sites = [Site(name=f"SITE {tag} {suffix}", client_id=client_row.id, active=1) for suffix in ("NORD", "SUD")]
    db.add_all(sites)
    db.commit()
    return {"tag": tag, "society": society, "position": position.name, "client": client_row, "sites": sites, "seed": seed}


def workbook(rows, titles=TITLES):
    book = Workbook()
    sheet = book.active
    sheet.append(titles)
    for row in rows:
        sheet.append(row)
    content = BytesIO()
    book.save(content)
    return content.getvalue()


def upload(client, headers, rows, titles=TITLES, society=None, name="employes.xlsx", expect=200):
    params = {"society": society} if society else {}
    response = client.post(BASE, params=params, headers=headers, files={"file": (name, workbook(rows, titles))})
    assert response.status_code == expect, response.text
    return response.json()


def mapping_of(result, sheet=0):
    return {str(column["index"]): column["suggested"] for column in result["sheets"][sheet]["columns"]}


def preview(client, headers, uploaded, mapping=None, expect=200):
    response = client.post(f"{BASE}/{uploaded['session_id']}/preview", headers=headers,
                           json={"sheet": 0, "mapping": mapping if mapping is not None else mapping_of(uploaded)})
    assert response.status_code == expect, response.text
    return response.json()


def correct(client, headers, sid, row, **payload):
    response = client.post(f"{BASE}/{sid}/corrections", headers=headers, json={"row": row, **payload})
    assert response.status_code == 200, response.text
    return response.json()


def confirm(client, headers, sid):
    response = client.post(f"{BASE}/{sid}/confirm", headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def row_of(result, number):
    return next(row for row in result["rows"] if row["row"] == number)


def accept_keys(row, field):
    return [issue["accept"] for issue in row["issues"] if issue["field"] == field and issue.get("accept")]


# ── Séparation Nom / Prénom (fonction pure) ──────────────────────────────────────────────────
@pytest.mark.parametrize("value,last,first,confidence", [
    ("BENALI Ahmed", "BENALI", "Ahmed", "certain"),
    ("BOUAZIZ Fatima Zahra", "BOUAZIZ", "Fatima Zahra", "certain"),              # prénom composé
    ("BOUAZIZ Fatima-Zahra", "BOUAZIZ", "Fatima-Zahra", "certain"),
    ("BEN ABDELKADER Sofiane", "BEN ABDELKADER", "Sofiane", "certain"),          # nom composé
    ("AIT EL HADJ Mohamed Amine", "AIT EL HADJ", "Mohamed Amine", "certain"),    # nom et prénom composés
    ("BENALI, Ahmed", "BENALI", "Ahmed", "certain"),
    ("BENALI AHMED", "BENALI", "AHMED", "review"),                               # casse uniforme
    ("benali ahmed", "benali", "ahmed", "review"),
    ("BEN ALI MOHAMED", "BEN ALI", "MOHAMED", "review"),
    ("Ahmed BENALI", "BENALI", "Ahmed", "review"),                               # ordre inversé
    ("BEN Abdelkader Sofiane", "BEN Abdelkader", "Sofiane", "review"),
    ("BENALI", "BENALI", "", "review"),
    ("BENALI 2 Ahmed", "BENALI", "2 Ahmed", "review"),
])
def test_split_full_name(value, last, first, confidence):
    result = employee_import.split_full_name(value)
    assert (result["last_name"], result["first_name"], result["confidence"]) == (last, first, confidence)
    assert result["original"] == value


# ── Parcours complet ─────────────────────────────────────────────────────────────────────────
def test_valid_file_creates_employees_contracts_and_audit(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    titles = ["Matricule", "Nom", "Prénom", "Téléphone", "Société", "Client", "Lieu", "Fonction", "Groupe",
              "Date de recrutement", "Type de contrat", "Date fin contrat"]
    rows = [[f"X{tag}1", f"ALPHA{tag}", "Karim", phone(), society, ref["client"].name, ref["sites"][0].name, ref["position"], "B",
             "01/10/2026", "CDD", "30/09/2027"],
            ["", f"BETA{tag}", "Nadia", phone(), society, "", "", ref["position"], "Groupe C", date(2026, 10, 5), "CDI", ""]]
    uploaded = upload(client, system_headers, rows, titles, society=society)
    suggested = mapping_of(uploaded)
    assert suggested == {"0": "code", "1": "last_name", "2": "first_name", "3": "phone", "4": "society", "5": "client", "6": "site",
                         "7": "position", "8": "group", "9": "recruit_date", "10": "contract_type", "11": "contract_end"}
    result = preview(client, system_headers, uploaded)
    assert result["counts"]["ready"] == 2, result
    assert not db.scalars(select(Employee).where(Employee.last_name.in_([f"ALPHA{tag}", f"BETA{tag}"]))).all()   # rien avant la confirmation
    outcome = confirm(client, system_headers, uploaded["session_id"])
    assert outcome["counts"]["created"] == 2 and outcome["counts"]["read"] == 2, outcome
    db.expire_all()
    alpha = db.scalar(select(Employee).where(Employee.code == f"X{tag}1"))
    assert alpha.last_name == f"ALPHA{tag}" and alpha.first_name == "KARIM" and alpha.position == ref["position"].upper()
    assert alpha.contract_type == "CDD" and alpha.contract_end_date == date(2027, 9, 30) and alpha.recruit_date == date(2026, 10, 1)
    assert alpha.extra["ficheACompleter"] is True and alpha.extra["recruitmentIncompleteFields"] == ["Date de naissance"]
    assert alpha.extra["importOrigin"]["propositionOps"] == {"groupe": "B", "statut": "à traiter par OPS"}
    contract = db.scalar(select(Contract).where(Contract.employee_id == alpha.id))
    assert contract.dc_client_id == ref["client"].id and contract.dc_site_id == ref["sites"][0].id
    assert contract.start_date == date(2026, 10, 1) and contract.trial_end_date == date(2026, 10, 31)
    # Le groupe n'est jamais une affectation OPS.
    assert not db.scalars(select(Assignment).where(Assignment.employee_id == alpha.id)).all()
    beta = db.scalar(select(Employee).where(Employee.last_name == f"BETA{tag}"))
    assert beta.code and beta.code != f"X{tag}1" and beta.contract_type == "CDI" and beta.contract_end_date is None
    audits = db.scalars(select(AuditEvent).where(AuditEvent.action == "drh.employees.import.created",
                                                 AuditEvent.resource_id.in_([str(alpha.id), str(beta.id)]))).all()
    assert len(audits) == 2
    assert db.scalar(select(AuditEvent).where(AuditEvent.action == "drh.employees.import",
                                              AuditEvent.resource_id == uploaded["session_id"]))
    # Idempotence : une seconde confirmation ne crée rien.
    again = confirm(client, system_headers, uploaded["session_id"])
    assert again["already_processed"] is True and again["counts"] == outcome["counts"]
    session = db.scalar(select(CandidateImportSession).where(CandidateImportSession.public_id == uploaded["session_id"]))
    assert session.workbook is None and session.status == "completed"


def test_full_name_in_one_cell_keeps_original_for_audit(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    titles = ["Nom salarié", "Téléphone"]
    rows = [[f"BOUAZIZ{tag} Fatima Zahra", phone()], [f"BEN ABDELKADER{tag} Sofiane", phone()]]
    uploaded = upload(client, system_headers, rows, titles, society=society)
    assert mapping_of(uploaded) == {"0": "full_name", "1": "phone"}
    result = preview(client, system_headers, uploaded)
    first, second = row_of(result, 2), row_of(result, 3)
    assert first["status"] == second["status"] == "ready", result
    assert first["name"]["detected"] == {"last_name": f"BOUAZIZ{tag}", "first_name": "Fatima Zahra"}
    assert second["name"]["detected"] == {"last_name": f"BEN ABDELKADER{tag}", "first_name": "Sofiane"}
    assert first["name"]["original"] == f"BOUAZIZ{tag} Fatima Zahra"
    outcome = confirm(client, system_headers, uploaded["session_id"])
    assert outcome["counts"]["created"] == 2 and outcome["counts"]["incomplete"] == 2
    db.expire_all()
    employee = db.scalar(select(Employee).where(Employee.last_name == f"BEN ABDELKADER{tag}"))
    assert employee.first_name == "SOFIANE"
    origin = employee.extra["importOrigin"]
    assert origin["nomCompletOriginal"] == f"BEN ABDELKADER{tag} Sofiane"
    assert origin["separationNom"] == {"nom": f"BEN ABDELKADER{tag}", "prenom": "Sofiane", "source": "auto", "confiance": "certain"}
    audit = db.scalar(select(AuditEvent).where(AuditEvent.action == "drh.employees.import.created", AuditEvent.resource_id == str(employee.id)))
    assert f"BEN ABDELKADER{tag} Sofiane" in audit.new_state
    assert outcome["rows"][1]["original_name"] == f"BEN ABDELKADER{tag} Sofiane"


def test_uncertain_name_is_blocked_until_corrected_or_validated(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    rows = [[f"GAMMA{tag} AMINE", phone()], [f"Yacine DELTA{tag}", phone()]]
    uploaded = upload(client, system_headers, rows, ["Nom complet", "Téléphone"], society=society)
    result = preview(client, system_headers, uploaded)
    assert row_of(result, 2)["status"] == row_of(result, 3)["status"] == "review"
    sid = uploaded["session_id"]
    # Ligne 2 : la proposition est validée explicitement ; ligne 3 : corrigée à la main.
    result = correct(client, system_headers, sid, 2, accept=accept_keys(row_of(result, 2), "name"))
    assert row_of(result, 2)["status"] == "ready" and row_of(result, 2)["name"]["source"] == "validated"
    result = correct(client, system_headers, sid, 3, values={"last_name": f"DELTA{tag}", "first_name": "Yacine"})
    assert row_of(result, 3)["status"] == "ready" and row_of(result, 3)["name"]["source"] == "manual"
    assert confirm(client, system_headers, sid)["counts"]["created"] == 2
    db.expire_all()
    corrected = db.scalar(select(Employee).where(Employee.last_name == f"DELTA{tag}"))
    assert corrected.first_name == "YACINE"
    assert corrected.extra["importOrigin"]["nomCompletOriginal"] == f"Yacine DELTA{tag}"
    assert corrected.extra["importOrigin"]["separationNom"]["source"] == "manual"


def test_unknown_column_is_rejected_and_unmapped_columns_are_ignored(client, system_headers, ref):
    tag, society = ref["tag"], ref["society"]
    titles = ["Nom", "Prénom", "Couleur préférée"]
    uploaded = upload(client, system_headers, [[f"EPSILON{tag}", "Ali", "Bleu"]], titles, society=society)
    assert mapping_of(uploaded)["2"] is None
    response = client.post(f"{BASE}/{uploaded['session_id']}/preview", headers=system_headers,
                           json={"sheet": 0, "mapping": {"0": "last_name", "1": "first_name", "9": "phone"}})
    assert response.status_code == 422 and "colonne inconnue" in response.text
    response = client.post(f"{BASE}/{uploaded['session_id']}/preview", headers=system_headers,
                           json={"sheet": 0, "mapping": {"0": "last_name", "1": "first_name", "2": "salaire_secret"}})
    assert response.status_code == 422 and "champ inconnu" in response.text
    response = client.post(f"{BASE}/{uploaded['session_id']}/preview", headers=system_headers,
                           json={"sheet": 0, "mapping": {"0": "full_name", "1": "first_name"}})
    assert response.status_code == 422 and "Nom complet" in response.text
    result = preview(client, system_headers, uploaded)
    assert result["ignored_columns"] == ["C (Couleur préférée)"]


def test_duplicate_matricule_in_database_and_in_file(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    rows = [[ref["seed"].code, f"ZETA{tag}", "Omar", phone()],
            [f"D{tag}", f"ETA{tag}", "Lina", phone()],
            [f"d{tag}", f"THETA{tag}", "Sami", phone()],
            ["", f"IOTA{tag}", "Rym", phone()]]
    uploaded = upload(client, system_headers, rows, TITLES[:4], society=society)
    result = preview(client, system_headers, uploaded)
    assert row_of(result, 2)["status"] == "blocked" and "déjà utilisé" in row_of(result, 2)["issues"][0]["message"]
    assert row_of(result, 3)["status"] == row_of(result, 4)["status"] == "blocked"
    assert "plusieurs fois dans le fichier" in str(row_of(result, 3)["issues"])
    assert row_of(result, 5)["status"] == "ready"
    outcome = confirm(client, system_headers, uploaded["session_id"])
    assert outcome["counts"]["created"] == 1 and outcome["counts"]["duplicate"] == 3, outcome
    db.expire_all()
    assert db.scalar(select(Employee).where(Employee.last_name == f"ZETA{tag}")) is None


def test_existing_employee_and_possible_duplicate(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    nin = "".join(random.choices(string.digits, k=18))
    existing = Employee(code=f"E{tag}", first_name="MOURAD", last_name=f"KAPPA{tag}", society=society, nin=nin, phone=phone())
    db.add(existing)
    db.commit()
    rows = [[f"LAMBDA{tag}", "Imane", phone(), nin], [f"KAPPA{tag}", "Mourad", phone(), ""]]
    uploaded = upload(client, system_headers, rows, ["Nom", "Prénom", "Téléphone", "NIN"], society=society)
    result = preview(client, system_headers, uploaded)
    assert row_of(result, 2)["status"] == "blocked" and "même NIN" in str(row_of(result, 2)["issues"])
    homonym = row_of(result, 3)
    assert homonym["status"] == "review" and "Doublon possible" in str(homonym["issues"])
    result = correct(client, system_headers, uploaded["session_id"], 3, accept=accept_keys(homonym, "duplicate"))
    assert row_of(result, 3)["status"] == "ready"
    outcome = confirm(client, system_headers, uploaded["session_id"])
    assert outcome["counts"]["created"] == 1 and outcome["counts"]["duplicate"] == 1


def test_unknown_site_is_proposed_then_completed_or_left_to_complete(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    titles = ["Nom", "Prénom", "Site", "Type de contrat", "Date début contrat"]
    rows = [[f"MU{tag}", "Adel", f"SITE {tag} NORDD", "CDD", "01/11/2026"],
            [f"NU{tag}", "Sara", "SITE INEXISTANT", "CDD", "01/11/2026"]]
    uploaded = upload(client, system_headers, rows, titles, society=society)
    result = preview(client, system_headers, uploaded)
    misspelled, unknown = row_of(result, 2), row_of(result, 3)
    assert misspelled["status"] == unknown["status"] == "review"
    issue = next(issue for issue in misspelled["issues"] if issue["field"] == "client_site")
    assert ref["sites"][0].name in issue["suggestions"][0]
    sid = uploaded["session_id"]
    result = correct(client, system_headers, sid, 2, values={"site": ref["sites"][0].name})
    assert row_of(result, 2)["status"] == "ready" and row_of(result, 2)["client_site"].endswith(ref["sites"][0].name)
    result = correct(client, system_headers, sid, 3, accept=accept_keys(row_of(result, 3), "client_site"))
    assert row_of(result, 3)["status"] == "ready" and "Client / Site" in row_of(result, 3)["incomplete"]
    assert confirm(client, system_headers, sid)["counts"]["created"] == 2
    db.expire_all()
    linked = db.scalar(select(Employee).where(Employee.last_name == f"MU{tag}"))
    assert db.scalar(select(Contract).where(Contract.employee_id == linked.id)).dc_site_id == ref["sites"][0].id
    to_complete = db.scalar(select(Employee).where(Employee.last_name == f"NU{tag}"))
    contract = db.scalar(select(Contract).where(Contract.employee_id == to_complete.id))
    assert contract.dc_client_id is None and contract.dc_site_id is None
    assert "Client / Site" in to_complete.extra["recruitmentIncompleteFields"]
    assert db.scalar(select(Site).where(Site.name == "SITE INEXISTANT")) is None       # jamais créé


def test_unknown_position_is_never_created(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    rows = [[f"XI{tag}", "Walid", "Plongeur sous-marin"], [f"OMICRON{tag}", "Hana", f"agent {tag.lower()}"]]
    uploaded = upload(client, system_headers, rows, ["Nom", "Prénom", "Fonction"], society=society)
    result = preview(client, system_headers, uploaded)
    assert row_of(result, 2)["status"] == "review" and "Poste inconnu" in str(row_of(result, 2)["issues"])
    assert row_of(result, 3)["status"] == "ready" and row_of(result, 3)["values"]["position"] == ref["position"]
    result = correct(client, system_headers, uploaded["session_id"], 2, accept=accept_keys(row_of(result, 2), "position"))
    assert row_of(result, 2)["status"] == "ready"
    confirm(client, system_headers, uploaded["session_id"])
    db.expire_all()
    employee = db.scalar(select(Employee).where(Employee.last_name == f"XI{tag}"))
    assert employee.position is None and employee.extra["ficheACompleter"] is True
    assert db.scalar(select(Position).where(Position.name == "Plongeur sous-marin")) is None


def test_incomplete_and_invalid_lines(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    titles = ["Nom", "Prénom", "Téléphone", "Date de recrutement", "Groupe", "Date de naissance"]
    rows = [[f"PI{tag}", "", phone(), "01/10/2026", "", ""],                 # prénom absent : bloquant
            [f"RHO{tag}", "Malik", "", "", "", ""],                         # facultatifs absents : fiche à compléter
            [f"SIGMA{tag}", "Kenza", phone(), "31/02/2026", "", ""],        # date critique invalide : bloquant
            [f"TAU{tag}", "Anis", phone(), "01/10/2026", "E", "12/12/1990"]]  # groupe invalide : à vérifier
    uploaded = upload(client, system_headers, rows, titles, society=society)
    result = preview(client, system_headers, uploaded)
    assert row_of(result, 2)["status"] == "blocked" and "Prénom obligatoire" in str(row_of(result, 2)["issues"])
    assert row_of(result, 3)["status"] == "ready"
    assert row_of(result, 3)["incomplete"] == ["Téléphone", "Poste", "Date de naissance", "Date de recrutement"]
    assert row_of(result, 4)["status"] == "blocked" and "n'existe pas" in str(row_of(result, 4)["issues"])
    assert row_of(result, 5)["status"] == "review" and "Groupe invalide" in str(row_of(result, 5)["issues"])
    assert result["counts"] == {"total": 4, "ready": 1, "review": 1, "blocked": 2, "excluded": 0, "incomplete": 1, "empty_rows": 0}


def test_partial_import_with_corrections_and_exclusion(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    rows = [[f"UPSILON{tag}", "Rafik", phone()], [f"PHI{tag}", "", phone()], [f"CHI{tag}", "Lamia", "12"], [f"PSI{tag}", "", phone()]]
    uploaded = upload(client, system_headers, rows, ["Nom", "Prénom", "Téléphone"], society=society)
    sid = uploaded["session_id"]
    result = preview(client, system_headers, uploaded)
    assert [row_of(result, n)["status"] for n in (2, 3, 4, 5)] == ["ready", "blocked", "review", "blocked"]
    result = correct(client, system_headers, sid, 3, values={"first_name": "Sofia"})   # correction directe dans l'aperçu
    result = correct(client, system_headers, sid, 5, exclude=True)
    assert row_of(result, 3)["status"] == "ready" and row_of(result, 5)["status"] == "excluded"
    outcome = confirm(client, system_headers, sid)
    assert outcome["counts"] == {"read": 4, "created": 2, "incomplete": 2, "duplicate": 0, "refused": 1, "excluded": 1, "error": 0,
                                 "empty_rows": 0}, outcome
    db.expire_all()
    assert db.scalar(select(Employee).where(Employee.last_name == f"PHI{tag}")).first_name == "SOFIA"
    assert db.scalar(select(Employee).where(Employee.last_name == f"CHI{tag}")) is None
    report = client.get(f"{BASE}/{sid}/report", headers=system_headers)
    assert report.status_code == 200
    sheet = load_workbook(BytesIO(report.content)).active
    values = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    assert ["Lignes lues", 4] in [list(row[:2]) for row in values]
    assert any(row[0] == 4 and row[1] == "Refusé" for row in values)


def test_row_changed_after_verification_is_not_created(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    uploaded = upload(client, system_headers, [[f"C{tag}", f"OMEGA{tag}", "Ilyes", phone()]], TITLES[:4], society=society)
    assert preview(client, system_headers, uploaded)["counts"]["ready"] == 1
    db.add(Employee(code=f"C{tag}", first_name="AUTRE", last_name=f"AUTRE{tag}", society=society))
    db.commit()
    outcome = confirm(client, system_headers, uploaded["session_id"])
    assert outcome["counts"]["created"] == 0 and outcome["counts"]["duplicate"] == 1
    assert "déjà utilisé" in outcome["rows"][0]["reason"]


# ── Sécurité : permissions, multi-sociétés, fichiers ─────────────────────────────────────────
def test_permissions(client, auth_headers, restricted_headers, system_headers, db):
    # Admin sans session « Administration système », utilisateur limité : refusés et audités.
    for headers in (auth_headers, restricted_headers):
        assert client.get(f"{BASE}/config", headers=headers).status_code == 403
        assert client.post(BASE, headers=headers, files={"file": ("e.xlsx", workbook([["A", "B"]]))}).status_code == 403
    assert db.scalar(select(AuditEvent).where(AuditEvent.action == "authorization.drh.employees.import"))
    admin = db.scalar(select(User).where(User.username == "testadmin"))
    previous = admin.authorized_actions
    try:
        admin.authorized_actions = ["read"]
        db.commit()
        response = client.get(f"{BASE}/config", headers=system_headers)
        assert response.status_code == 403 and "autorisés à créer" in response.text
    finally:
        admin.authorized_actions = previous
        db.commit()
    assert client.get(f"{BASE}/config", headers=system_headers).status_code == 200


def test_multi_society_isolation(client, system_headers, db, ref):
    tag, society = ref["tag"], ref["society"]
    other = f"SOC AUTRE {tag}"
    hidden = Employee(code=f"H{tag}", first_name="CACHE", last_name=f"HIDDEN{tag}", society=other, phone=phone())
    db.add(hidden)
    db.commit()
    admin = db.scalar(select(User).where(User.username == "testadmin"))
    previous = (admin.global_society_access, admin.authorized_societies)
    try:
        admin.global_society_access, admin.authorized_societies = False, [society]
        db.commit()
        # Société active hors périmètre : refusée dès l'envoi.
        response = client.post(BASE, params={"society": other}, headers=system_headers, files={"file": ("e.xlsx", workbook([["A", "B"]]))})
        assert response.status_code == 403
        titles = ["Nom", "Prénom", "Société", "Téléphone"]
        rows = [[f"ALEPH{tag}", "Rania", other, phone()], [f"BETH{tag}", "Fares", society, phone()],
                [f"HIDDEN{tag}", "Cache", society, phone()], [f"GIMEL{tag}", "Lyes", "SOCIETE FANTOME", phone()]]
        uploaded = upload(client, system_headers, rows, titles)
        result = preview(client, system_headers, uploaded)
        assert row_of(result, 2)["status"] == "blocked" and "hors de votre périmètre" in str(row_of(result, 2)["issues"])
        assert row_of(result, 3)["status"] == "ready"
        homonym = str(row_of(result, 4)["issues"])
        assert "hors de votre périmètre" in homonym and f"H{tag}" not in homonym      # aucune fuite d'une autre société
        assert row_of(result, 5)["status"] == "review" and "Société inconnue" in str(row_of(result, 5)["issues"])
        # Société active choisie : toute autre société dans le fichier est bloquée.
        uploaded = upload(client, system_headers, rows[:2], titles, society=society)
        result = preview(client, system_headers, uploaded)
        assert "différente de la société active" in str(row_of(result, 2)["issues"])
    finally:
        admin.global_society_access, admin.authorized_societies = previous
        db.commit()


@pytest.mark.parametrize("name,content,message", [
    ("employes.csv", b"Nom;Prenom\nA;B", "Format non pris en charge"),
    ("employes.xlsx", b"ceci n'est pas un classeur", "Format incorrect"),
    ("employes.xlsx", b"", "vide"),
    ("employes.xlsx", b"PK\x03\x04corrompu", "corrompu"),
])
def test_invalid_file_is_refused_and_audited(client, system_headers, db, name, content, message):
    response = client.post(BASE, headers=system_headers, files={"file": (name, content)})
    assert response.status_code == 422 and message in response.text
    assert db.scalar(select(AuditEvent).where(AuditEvent.action == "drh.employees.import.analyze", AuditEvent.result == "refused"))


def test_sessions_are_isolated_from_candidate_imports(client, system_headers, db, ref):
    candidate = client.post("/api/drh/candidates/import", headers=system_headers,
                            files={"file": ("c.xlsx", workbook([["DUPONT", "Jean"]], ["Nom", "Prénom"]))})
    assert candidate.status_code == 200, candidate.text
    employees = upload(client, system_headers, [[f"ISO{ref['tag']}", "Ali"]], ["Nom", "Prénom"], society=ref["society"])
    db.expire_all()
    status = db.scalar(select(CandidateImportSession.status).where(CandidateImportSession.public_id == candidate.json()["session_id"]))
    assert status == "uploaded"                         # l'analyse candidats n'est pas annulée
    assert client.post(f"/api/drh/candidates/import/{employees['session_id']}/cancel", headers=system_headers).status_code == 404
    assert client.post(f"{BASE}/{candidate.json()['session_id']}/cancel", headers=system_headers).status_code == 404
    assert candidate_import.session_purpose(employees["session_id"]) == "employees"


def test_template_lists_existing_references(client, system_headers, ref):
    response = client.get(f"{BASE}/template", params={"society": ref["society"]}, headers=system_headers)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Employés", "Listes", "Instructions"]
    headers = [cell.value for cell in book["Employés"][1]]
    assert headers[:3] == ["Matricule", "Nom *", "Prénom *"]
    listed = [cell.value for row in book["Listes"].iter_rows() for cell in row if cell.value]
    assert ref["society"] in listed and ref["position"] in listed and "D" in listed
    # Le modèle rempli est relu par l'import (feuilles d'aide ignorées).
    describe = employee_import.describe(candidate_import.read_workbook("m.xlsx", response.content))
    assert [sheet["importable"] for sheet in describe["sheets"]] == [False, False, False]   # aucune ligne de données


def test_direct_employee_route_uses_shared_creation_and_audit(client, auth_headers, db, ref):
    tag = ref["tag"]
    response = client.post("/api/drh/employees", headers=auth_headers, json={
        "code": f"R{tag}", "first_name": "Route", "last_name": f"ROUTE{tag}", "society": ref["society"]})
    assert response.status_code == 200, response.text
    assert response.json()["code"] == f"R{tag}"
    assert db.scalar(select(AuditEvent).where(AuditEvent.action == "drh.employee.created",
                                              AuditEvent.resource_id == str(response.json()["id"])))
    # Code déjà pris : la fiche DRH garde son comportement historique (prochain code libre).
    again = client.post("/api/drh/employees", headers=auth_headers, json={
        "code": f"R{tag}", "first_name": "Route", "last_name": f"ROUTEBIS{tag}", "society": ref["society"]})
    assert again.status_code == 200 and again.json()["code"] != f"R{tag}"
