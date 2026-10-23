"""Import Excel des employés en masse (Administration système › Recrutement groupé).

Même moteur que l'import des candidats (candidate_import.py) : lecture du classeur côté
serveur, sessions `employees-…` de la table candidate_import_sessions, correspondance des
colonnes, lecture des cellules, neutralisation des formules. Ce module n'ajoute que les règles
propres à la fiche employé :

- Nom complet dans une seule cellule → Nom et Prénom proposés, jamais figés sans vérification ;
- société, client, site, poste et groupe rapprochés des données existantes, jamais créés ;
- statut de ligne : `ready` (vert), `review` (orange : à compléter / à vérifier), `blocked` (rouge) ;
- corrections et validations explicites enregistrées dans la session, puis tout est revérifié ;
- confirmation ligne par ligne (import partiel) par `service.create_employee_direct`, la fonction
  unique de création d'employé, avec un audit par employé et un audit global.

Client et site ne sont écrits que sur le contrat initial (identifiants DC). Le groupe A/B/C/D
n'est jamais une affectation : il reste une proposition tracée dans l'origine de l'import ;
OPS reste l'autorité d'affectation."""
from __future__ import annotations

import difflib
import hashlib
import io
import json
import logging
import re
import secrets
from dataclasses import replace
from datetime import date, datetime, timedelta
from typing import Any

from fastapi import HTTPException
from sqlalchemy import select, text, update
from sqlalchemy.orm import Session

from app.core.audit import append_audit
from app.core.scope_policy import ScopeKind, authorized_society_values, society_key, society_scope
from app.modules.drh import candidate_import as excel, dc_client_sites, service
from app.modules.drh import employee_code_policy as code_policy
from app.modules.drh.candidate_import_models import CandidateImportSession
from app.modules.drh.candidate_import_reference import FIELD_BY_KEY as CANDIDATE_FIELDS, ImportField, norm_key
from app.modules.drh.models import Employee
from app.modules.drh.schemas import EmployeeCreate
from app.modules.irongs.constants import SOCIETES
from app.modules.irongs.models import Position

logger = logging.getLogger("sgdi.drh.employee_import")

PURPOSE = "employees"
PREFIX = excel.SESSION_PREFIXES[PURPOSE]
_ADVISORY_LOCK_KEY = 74202610230001
CONTRACT_TYPES = ("CDD", "CDI", "CIDD", "CTA", "APPRENTISSAGE")
TRIAL_DAYS = {"CDI": 90, "CDD": 30, "CIDD": 45}           # mêmes durées que service.recruit_candidate
GROUPS = ("A", "B", "C", "D")
_CODE_PATTERN = re.compile(r"^[A-Z0-9][A-Z0-9._/-]{0,29}$")
# Particules d'un nom de famille : « BEN ABDELKADER », « AIT EL HADJ », « OULD … ».
PARTICLES = frozenset({"ben", "bent", "bou", "ait", "at", "id", "ou", "ould", "ouled", "el", "al", "abd", "abdel",
                       "si", "sidi", "moulay", "de", "du", "des", "di", "da", "le", "la", "bel", "beni", "ba",
                       "بن", "ايت", "ولد"})


def _reuse(key: str, **changes: Any) -> ImportField:
    """Champ identique à celui de l'import candidats (mêmes alias et règles de lecture)."""
    return replace(CANDIDATE_FIELDS[key], column=None, data_key=None, required=False, **changes)


FIELDS: tuple[ImportField, ...] = (
    ImportField("code", "Matricule", "matricule", max_len=30, text_format=True, width=14, example="A101",
                aliases=("matricule", "mat", "n matricule", "numero matricule", "code", "code employe", "code agent", "id employe"),
                hint="Vide : attribué par ATLAS selon la série de la société"),
    ImportField("full_name", "Nom complet (Nom + Prénom)", "full_name", max_len=200, width=30, example="BENALI Ahmed",
                aliases=("nom complet", "nom et prenom", "nom et prenoms", "nom prenom", "nom prenoms", "nom salarie",
                         "nom du salarie", "salarie", "employe", "nom employe", "agent", "nom de l agent", "nom agent", "full name",
                         "الاسم واللقب", "اللقب والاسم"),
                hint="NOM en majuscules puis prénom(s) : BENALI Ahmed"),
    _reuse("last_name", aliases=("nom", "nom de famille", "last name", "surname", "اللقب", "لقب")),
    _reuse("first_name", aliases=("prenom", "prenoms", "first name", "الاسم", "اسم")),
    _reuse("phone", aliases=CANDIDATE_FIELDS["phone"].aliases + ("gsm", "mobile", "portable")),
    ImportField("society", "Société", "society", max_len=150, width=24,
                aliases=("societe", "entreprise", "filiale", "employeur", "company", "structure employeur")),
    ImportField("client", "Client", "client", max_len=200, width=24, aliases=("client", "donneur d ordre", "nom client")),
    ImportField("site", "Site", "site", max_len=200, width=24,
                aliases=("site", "lieu", "lieu de travail", "site d affectation", "site affectation", "affectation", "nom du site")),
    ImportField("position", "Poste", "position_ref", max_len=150, width=24,
                aliases=("poste", "fonction", "emploi", "qualification", "poste occupe", "poste de travail", "intitule du poste")),
    ImportField("group", "Groupe", "group", max_len=20, width=10, example="A", hint="A, B, C ou D",
                aliases=("groupe", "grp", "equipe", "groupe de travail", "groupe rotation")),
    ImportField("recruit_date", "Date de recrutement", "date", width=18, example="01/10/2026", hint="JJ/MM/AAAA",
                aliases=("date de recrutement", "date recrutement", "date d embauche", "date embauche", "date d entree",
                         "date entree", "recrute le", "embauche le")),
    ImportField("contract_type", "Type de contrat", "choice", width=14, example="CDD",
                choices=(("CDD", ("cdd", "duree determinee", "contrat a duree determinee")),
                         ("CDI", ("cdi", "duree indeterminee", "contrat a duree indeterminee")),
                         ("CIDD", ("cidd",)), ("CTA", ("cta",)), ("APPRENTISSAGE", ("apprentissage", "apprenti"))),
                aliases=("type de contrat", "type contrat", "contrat", "nature du contrat", "nature contrat")),
    ImportField("contract_start", "Date début contrat", "date", width=18, hint="JJ/MM/AAAA",
                aliases=("date debut contrat", "debut contrat", "date de debut", "date debut", "debut du contrat")),
    ImportField("contract_end", "Date fin contrat", "date", width=18, hint="JJ/MM/AAAA",
                aliases=("date fin contrat", "fin contrat", "date de fin", "date fin", "fin du contrat", "echeance", "echeance contrat")),
    _reuse("nin"),
    _reuse("birth_date"),
    _reuse("birth_place"),
    _reuse("father_name"),
    _reuse("mother_name"),
    _reuse("family_status"),
    _reuse("children_count"),
    _reuse("email"),
    _reuse("address"),
    _reuse("wilaya"),
    _reuse("commune"),
    ImportField("salary_net", "Salaire net", "money", width=14, aliases=("salaire", "salaire net", "salaire mensuel", "net a payer")),
)
FIELD_BY_KEY: dict[str, ImportField] = {item.key: item for item in FIELDS}
# Une erreur de lecture sur ces champs bloque la ligne ; ailleurs, elle peut être corrigée ou écartée.
CRITICAL_FIELDS = frozenset({"code", "full_name", "last_name", "first_name", "society", "recruit_date", "contract_start", "contract_end"})
# Même liste que l'écran Recrutement groupé (adminRecruitmentMissingFields) : absentes → FICHE À COMPLÉTER.
COMPLETION_FIELDS = ("phone", "position", "birth_date", "recruit_date")
PSEUDO_LABELS = {"name": "Nom / Prénom", "duplicate": "Doublon", "client_site": "Client / Site"}


def _aliases() -> dict[str, str]:
    table: dict[str, str] = {}
    for item in FIELDS:
        for alias in (item.label, *item.aliases):
            table.setdefault(norm_key(alias), item.key)
    return table


_HEADER_ALIASES = _aliases()


def field_for_header(header: object) -> str | None:
    return _HEADER_ALIASES.get(norm_key(header))


def label(field: str) -> str:
    return FIELD_BY_KEY[field].label if field in FIELD_BY_KEY else PSEUDO_LABELS.get(field, field)


def describe(workbook: dict[str, Any]) -> dict[str, Any]:
    """Feuilles et correspondance proposée. Une colonne « Nom » sans colonne « Prénom » contient
    presque toujours le nom complet : elle est proposée comme Nom complet."""
    result = excel.describe_workbook(workbook, field_for_header)
    for sheet in result["sheets"]:
        suggested = {column["suggested"] for column in sheet["columns"]}
        if "last_name" in suggested and "first_name" not in suggested and "full_name" not in suggested:
            for column in sheet["columns"]:
                if column["suggested"] == "last_name":
                    column["suggested"] = "full_name"
    return result


# ── Nom complet → Nom + Prénom ──────────────────────────────────────────────────────────────
def _token_case(token: str) -> str:
    letters = [ch for ch in token if ch.isalpha()]
    if not letters:
        return "none"
    if not any(ch.isupper() or ch.islower() for ch in letters):
        return "uncased"                                   # écriture sans casse (arabe)
    if all(ch.isupper() for ch in letters if ch.isupper() or ch.islower()):
        return "upper" if len(letters) >= 2 else "initial"
    parts = [part for part in re.split(r"[-'’]", token) if part]
    if all(part[:1].isupper() and part[1:] == part[1:].lower() for part in parts):
        return "title"
    return "lower" if token == token.lower() else "mixed"


def _is_particle(token: str) -> bool:
    return norm_key(token) in PARTICLES


def split_full_name(value: Any) -> dict[str, Any]:
    """Propose Nom et Prénom à partir d'une cellule « Nom complet ». Rien n'est figé : la
    proposition est « certaine » seulement quand la casse sépare sans ambiguïté le NOM (en
    majuscules, en tête) du ou des prénoms ; sinon la ligne est « à vérifier »."""
    original = " ".join(str(value or "").split())

    def result(last: list[str] | str, first: list[str] | str, confidence: str, reason: str) -> dict[str, Any]:
        last = " ".join(last) if isinstance(last, list) else last
        first = " ".join(first) if isinstance(first, list) else first
        return {"original": original, "last_name": last.strip(), "first_name": first.strip(), "confidence": confidence, "reason": reason}

    if "," in original:
        before, after = (part.strip() for part in original.split(",", 1))
        if before and after:
            return result(before, after, "certain", "Séparateur virgule : NOM, Prénom")
    tokens = original.split(" ") if original else []
    if len(tokens) < 2:
        return result(tokens, "", "review", "Un seul mot : prénom introuvable")
    if any(ch.isdigit() for ch in original):
        return result(tokens[:1], tokens[1:], "review", "Chiffres dans le nom")
    cases = [_token_case(token) for token in tokens]
    mixed = {"title", "lower", "mixed"}
    if cases[0] == "upper" and any(case != "upper" for case in cases):
        split = next(index for index, case in enumerate(cases) if case != "upper")
        if all(case in mixed for case in cases[split:]):
            if all(_is_particle(token) for token in tokens[:split]):
                if split + 1 < len(tokens):
                    return result(tokens[:split + 1], tokens[split + 1:], "review", "Particule seule en majuscules : nom supposé")
                return result(tokens[:split], tokens[split:], "review", "Particule seule en majuscules : nom supposé")
            return result(tokens[:split], tokens[split:], "certain", "NOM en majuscules puis prénom")
        return result(tokens[:split], tokens[split:], "review", "Majuscules dispersées : séparation à vérifier")
    if cases[-1] == "upper" and cases[0] in mixed:
        split = next(index for index, case in enumerate(cases) if case == "upper")
        if all(case == "upper" for case in cases[split:]) and all(case in mixed for case in cases[:split]):
            return result(tokens[split:], tokens[:split], "review", "Ordre Prénom NOM détecté : à vérifier")
    # Casse uniforme : on suppose NOM puis prénom, particules rattachées au nom.
    index = 0
    while index < len(tokens) - 1 and _is_particle(tokens[index]):
        index += 1
    index += 1
    if index >= len(tokens):
        index = 1
    return result(tokens[:index], tokens[index:], "review", "Casse uniforme : séparation supposée NOM puis prénom")


# ── Contexte d'analyse : référentiels existants, construits une fois ───────────────────────
class _Context:
    def __init__(self, db: Session, user: Any, active_society: str | None):
        self.db, self.user, self.active = db, user, active_society
        self.scope = society_scope(user)
        self.societies: dict[str, str] = {}
        for value in [*SOCIETES, *authorized_society_values(user), *service.known_societies(db)]:
            label_value = str(value or "").strip()
            if label_value:
                self.societies.setdefault(society_key(label_value), label_value)
        self.allowed = {key: value for key, value in self.societies.items() if self.scope.allows(value)}
        # Aucun amorçage du référentiel ici : seuls les postes actifs existants sont proposés.
        self.positions = [(name, society) for name, society in
                          db.execute(select(Position.name, Position.society).where(Position.active.is_(True))).all()]
        self._options: dict[str, list[dict[str, Any]]] = {}
        self.codes: set[str] = set()
        self.by_nin: dict[str, list[Any]] = {}
        self.by_phone: dict[str, list[Any]] = {}
        self.by_name: dict[str, list[Any]] = {}
        columns = (Employee.id, Employee.code, Employee.first_name, Employee.last_name, Employee.phone, Employee.nin,
                   Employee.birth_date, Employee.society)
        for row in db.execute(select(*columns)).all():
            self.codes.add(str(row.code or "").strip().upper())
            if str(row.nin or "").strip():
                self.by_nin.setdefault(str(row.nin).strip(), []).append(row)
            phone = _phone_key(row.phone)
            if phone:
                self.by_phone.setdefault(phone, []).append(row)
            name = _name_key(row.last_name, row.first_name)
            if name:
                self.by_name.setdefault(name, []).append(row)

    def positions_for(self, society: str | None) -> dict[str, str]:
        wanted = society_key(society)
        reference: dict[str, str] = {}
        for name, owner in sorted(self.positions, key=lambda item: (item[1] is not None, item[0])):
            if owner is None or society_key(owner) == wanted:
                reference.setdefault(norm_key(name), name)
        return reference

    def client_sites(self, society: str) -> list[dict[str, Any]]:
        key = society_key(society)
        if key not in self._options:
            self._options[key] = dc_client_sites.list_client_site_options(self.db, society)
        return self._options[key]

    def visible(self, society: Any) -> bool:
        return self.scope.kind is ScopeKind.GLOBAL or self.scope.allows(society)


def _phone_key(value: Any) -> str:
    key = service._normalized_contact_phone(value)
    return key if len(key) >= 8 else ""


def _name_key(last: Any, first: Any) -> str:
    last_key, first_key = norm_key(last), norm_key(first)
    return f"{last_key}|{first_key}" if last_key and first_key else ""


def _suggest(value: str, labels: list[str], limit: int = 3) -> list[str]:
    by_key = {norm_key(item): item for item in labels}
    return [by_key[key] for key in difflib.get_close_matches(norm_key(value), list(by_key), n=limit, cutoff=0.6)]


ACCEPT_LABELS = {"name": "Valider la séparation proposée", "duplicate": "Créer malgré ce doublon possible"}
DROP_LABEL = "Valider sans cette valeur : la fiche sera à compléter"


def _issue(level: str, field: str, message: str, *, accept: str | None = None, suggestions: list[str] | None = None,
           accept_label: str | None = None) -> dict[str, Any]:
    issue = {"level": level, "field": field, "label": label(field), "message": message}
    if accept:
        # La validation explicite vise CE constat : un autre constat sur la même ligne en exige une nouvelle.
        issue["accept"] = f"{accept}:{hashlib.sha1(message.encode()).hexdigest()[:8]}"
        issue["accept_label"] = accept_label or ACCEPT_LABELS.get(accept, DROP_LABEL)
    if suggestions:
        issue["suggestions"] = suggestions
    return issue


def _display(cell: Any) -> str:
    return excel.cell_display(cell)


# ── Analyse d'une ligne ─────────────────────────────────────────────────────────────────────
def _analyze_row(ctx: _Context, row_number: int, cells: dict[str, Any], overrides: dict[str, Any],
                 accepted: set[str]) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    values: dict[str, Any] = {}
    dropped: list[str] = []                               # champs écartés par validation explicite
    effective = dict(cells)
    for field, value in overrides.items():
        if field in FIELD_BY_KEY:
            effective[field] = value if value not in (None, "") else None

    def drop_or_review(field: str, message: str, accept: str, suggestions: list[str] | None = None) -> None:
        issue = _issue("review", field, message, accept=accept, suggestions=suggestions)
        if issue["accept"] in accepted:
            dropped.append(field)
        else:
            issues.append(issue)

    # 1. Lecture des cellules, mêmes règles que l'import candidats.
    for item in FIELDS:
        if item.key not in effective:
            continue
        cell = effective[item.key]
        try:
            if isinstance(cell, dict) and "e" in cell:
                raise excel._FieldError(f"Cellule inexploitable ({cell['e']})")
            if not _display(cell):
                continue
            if item.kind in ("matricule", "full_name", "society", "client", "site", "position_ref", "group"):
                raw = excel._text_of(cell)
                if len(raw) > item.max_len:
                    raise excel._FieldError(f"{item.label} : texte trop long ({len(raw)} caractères, maximum {item.max_len})")
                values[item.key] = raw
                continue
            values[item.key], warning = excel.convert_cell(item, cell, values, {})
            if warning:
                issues.append(_issue("warning", item.key, warning))
        except excel._FieldError as exc:
            if item.key in CRITICAL_FIELDS:
                issues.append(_issue("blocked", item.key, str(exc)))
            else:
                drop_or_review(item.key, f"{exc} — corrigez la valeur ou validez pour l'ignorer", "invalid")

    # 2. Nom et prénom : cellule unique proposée, colonnes séparées ou correction manuelle.
    name: dict[str, Any] = {"original": None, "source": "columns", "confidence": "certain", "reason": ""}
    manual = {key: overrides.get(key) for key in ("last_name", "first_name") if overrides.get(key) not in (None, "")}
    if "full_name" in values:
        proposal = split_full_name(values["full_name"])
        name.update(original=proposal["original"], detected={"last_name": proposal["last_name"], "first_name": proposal["first_name"]},
                    source="auto", confidence=proposal["confidence"], reason=proposal["reason"])
        values.setdefault("last_name", proposal["last_name"])
        values.setdefault("first_name", proposal["first_name"])
        if manual:
            name.update(source="manual", confidence="certain", reason="Corrigé dans l'aperçu")
        elif proposal["confidence"] != "certain" and proposal["last_name"] and proposal["first_name"]:
            issue = _issue("review", "name", f"{proposal['reason']} : « {proposal['original']} » → NOM « {proposal['last_name']} », "
                           f"Prénom « {proposal['first_name']} »", accept="name")
            if issue["accept"] in accepted:
                name.update(source="validated")
            else:
                issues.append(issue)
    for key in ("last_name", "first_name"):
        if key in manual:
            values[key] = " ".join(str(manual[key]).split())[:100]
        if not str(values.get(key) or "").strip():
            if not any(issue["field"] == key for issue in issues):
                issues.append(_issue("blocked", key, f"{label(key)} obligatoire"))
            values.pop(key, None)
        elif len(values[key]) < 2:
            issues.append(_issue("blocked", key, f"{label(key)} : au moins 2 caractères"))
    name["last_name"], name["first_name"] = values.get("last_name", ""), values.get("first_name", "")

    # 3. Société : existante, dans le périmètre, et égale à la société active si elle est choisie.
    society = None
    raw_society = values.get("society") or ctx.active
    if not raw_society:
        if not any(issue["field"] == "society" for issue in issues):
            issues.append(_issue("blocked", "society", "Société obligatoire"))
    else:
        key = society_key(raw_society)
        if ctx.active and key != society_key(ctx.active):
            issues.append(_issue("blocked", "society", f"Société « {raw_society} » différente de la société active « {ctx.active} »"))
        elif key in ctx.allowed:
            society = ctx.allowed[key]
        elif key in ctx.societies:
            issues.append(_issue("blocked", "society", f"Société « {raw_society} » hors de votre périmètre"))
        else:
            issues.append(_issue("review", "society", f"Société inconnue : « {raw_society} » — choisissez une société existante",
                                 suggestions=_suggest(raw_society, list(ctx.allowed.values()))))
    values["society"] = society

    # 4. Poste : référentiel existant (postes communs ou propres à la société).
    if values.get("position"):
        positions = ctx.positions_for(society)
        resolved = positions.get(norm_key(values["position"]))
        if resolved:
            values["position"] = resolved
        else:
            drop_or_review("position", f"Poste inconnu : « {values['position']} »", "position",
                           _suggest(values["position"], list(positions.values())))
            values.pop("position")

    # 5. Groupe : proposition A/B/C/D pour OPS, jamais une affectation.
    if values.get("group"):
        group = re.sub(r"^(groupe|grp|gr|g)\s*", "", norm_key(values["group"])).upper()
        if group in GROUPS:
            values["group"] = group
        else:
            drop_or_review("group", f"Groupe invalide : « {values['group']} » (attendu A, B, C ou D)", "group")
            values.pop("group")

    # 6. Dates et type de contrat.
    start = values.get("contract_start") or values.get("recruit_date")
    values.setdefault("recruit_date", start)
    if not values.get("recruit_date"):
        values.pop("recruit_date", None)
    contract_type = values.get("contract_type")
    end = values.get("contract_end")
    if end and start and end < start:
        issues.append(_issue("blocked", "contract_end", "Date de fin du contrat antérieure à la date de début"))
    if end and contract_type == "CDI":
        issues.append(_issue("warning", "contract_end", "Contrat CDI : la date de fin est ignorée"))
        values.pop("contract_end")
    if (values.get("contract_start") or end) and not contract_type:
        drop_or_review("contract_type", "Dates de contrat sans type de contrat : contrat initial non créé", "contract_dates")

    # 7. Client / Site : structures officielles DC de la société, sur le contrat initial.
    link = None
    if (values.get("client") or values.get("site")) and society:
        options = ctx.client_sites(society)
        client_key, site_key = norm_key(values.get("client")), norm_key(values.get("site"))
        matches = [option for option in options
                   if (not client_key or norm_key(option["client_name"]) == client_key)
                   and (not site_key or norm_key(option["site_name"]) == site_key)]
        wanted = " — ".join(part for part in (values.get("client"), values.get("site")) if part)
        if len(matches) == 1:
            link = {"dc_client_id": matches[0]["dc_client_id"], "dc_site_id": matches[0]["dc_site_id"] if site_key else None,
                    "label": matches[0]["label"] if site_key else matches[0]["client_name"]}
        elif len({option["dc_client_id"] for option in matches}) == 1 and client_key and not site_key:
            link = {"dc_client_id": matches[0]["dc_client_id"], "dc_site_id": None, "label": matches[0]["client_name"]}
        else:
            reason = "ambigu (plusieurs correspondances)" if matches else "introuvable parmi les clients et sites actifs de la société"
            drop_or_review("client_site", f"Client / site « {wanted} » {reason}", "client_site",
                           [option["label"] for option in matches[:3]] or _suggest(wanted, [option["label"] for option in options]))
        if link and not contract_type:
            drop_or_review("client_site", "Client / site sans type de contrat : le rattachement exige le contrat initial", "client_site_contract")
            link = None if "client_site" in dropped else link
    values.pop("client", None)
    values.pop("site", None)
    if "client_site" in dropped:
        link = None

    # 8. Matricule : conservé s'il est libre, sinon ligne bloquée. Une forme non conforme à la
    # convention ATLAS (A001, A1) n'est jamais réécrite : l'utilisateur choisit explicitement.
    code_meta = None
    if values.get("code"):
        code = values["code"].strip().upper()
        values["code"] = code
        excel_code = _display(cells.get("code")).strip().upper() or None
        proposed = code_policy.canonical_code(excel_code) if excel_code else None
        if overrides.get("code") not in (None, ""):
            choice = "forme_proposee" if proposed and code == proposed and proposed != excel_code else "corrige_apercu"
        else:
            choice = "excel"
        if not _CODE_PATTERN.fullmatch(code):
            issues.append(_issue("blocked", "code", f"Matricule invalide : « {code} » (lettres, chiffres, . _ / -)"))
        elif code in ctx.codes:
            issues.append(_issue("blocked", "code", f"Matricule {code} déjà utilisé"))
        else:
            canonical = code_policy.canonical_code(code)
            if canonical != code:
                taken = canonical in ctx.codes
                issue = _issue("review", "code", f"Matricule Excel « {code} » ; forme proposée par ATLAS : « {canonical} »"
                               + (" (déjà utilisée par un autre employé)" if taken else ""),
                               accept="code_keep", accept_label=f"Conserver « {code} » tel quel comme matricule officiel",
                               suggestions=None if taken else [canonical])
                if issue["accept"] in accepted:
                    choice = "conserve_excel"
                else:
                    issues.append(issue)
            if society and code[:1] not in service.employee_code_prefixes_for_society(society):
                issues.append(_issue("warning", "code", f"Matricule {code} hors de la série de la société {society}"))
        code_meta = {"valeur": code, "valeurExcel": excel_code, "formeProposee": proposed if proposed and proposed != excel_code else None,
                     "choix": choice}

    # 9. Employé déjà existant (dans ATLAS).
    duplicate_reasons: list[str] = []
    suspicion: list[str] = []
    name_key = _name_key(values.get("last_name"), values.get("first_name"))
    phone_key = _phone_key(values.get("phone"))

    def who(row: Any) -> str:
        return f"{row.code} {row.last_name} {row.first_name}".strip() if ctx.visible(row.society) else "une fiche hors de votre périmètre"

    for row in ctx.by_nin.get(str(values.get("nin") or ""), []) if values.get("nin") else []:
        duplicate_reasons.append(f"même NIN que {who(row)}")
    for row in ctx.by_phone.get(phone_key, []) if phone_key else []:
        if name_key and _name_key(row.last_name, row.first_name) == name_key:
            duplicate_reasons.append(f"même nom et même téléphone que {who(row)}")
        else:
            suspicion.append(f"même téléphone que {who(row)}")
    for row in ctx.by_name.get(name_key, []) if name_key else []:
        if values.get("birth_date") and row.birth_date and str(row.birth_date) == str(values["birth_date"]):
            duplicate_reasons.append(f"même nom et même date de naissance que {who(row)}")
        elif not phone_key or _phone_key(row.phone) != phone_key:
            suspicion.append(f"même nom que {who(row)}")
    for reason in dict.fromkeys(duplicate_reasons):
        issues.append(_issue("blocked", "duplicate", f"Employé déjà existant : {reason}"))
    if suspicion and not duplicate_reasons:
        message = "Doublon possible : " + " ; ".join(dict.fromkeys(suspicion))
        issue = _issue("review", "duplicate", message, accept="duplicate")
        if issue["accept"] not in accepted:
            issues.append(issue)

    for field in dropped:
        values.pop(field, None)
    incomplete = [label(field) for field in COMPLETION_FIELDS if not values.get(field)]
    incomplete += [label(field) for field in dropped if label(field) not in incomplete]
    return {"row": row_number, "values": values, "link": link, "name": name, "code_meta": code_meta, "issues": issues, "dropped": dropped,
            "incomplete": incomplete, "cells": {key: _display(cell) for key, cell in effective.items()},
            "duplicate_keys": {"nin": str(values.get("nin") or ""), "identity": f"{name_key}|{phone_key}" if name_key and phone_key else "",
                               "name": name_key}}


def _status(row: dict[str, Any]) -> str:
    if row.get("excluded"):
        return "excluded"
    levels = {issue["level"] for issue in row["issues"]}
    return "blocked" if "blocked" in levels else "review" if "review" in levels else "ready"


def _employee_payload(row: dict[str, Any]) -> dict[str, Any]:
    """Valeurs de création de l'employé et du contrat initial : seule base de l'empreinte de
    vérification (une ligne qui change après la vérification n'est jamais créée)."""
    values = row["values"]
    start = values.get("contract_start") or values.get("recruit_date")
    contract_type = values.get("contract_type")
    trial_end = None
    if contract_type and start:
        trial_end = (date.fromisoformat(start) + timedelta(days=TRIAL_DAYS.get(contract_type, 30))).isoformat()
    recruit = values.get("recruit_date")
    employee = {key: values.get(key) for key in ("code", "last_name", "first_name", "phone", "society", "position", "nin", "birth_date",
                                                 "birth_place", "father_name", "mother_name", "family_status", "children_count",
                                                 "email", "address", "commune", "wilaya", "salary_net")}
    employee.update(recruit_date=recruit, contract_type=contract_type, trial_end_date=trial_end,
                    contract_end_date=values.get("contract_end") if contract_type else None,
                    status="a_venir" if recruit and recruit > date.today().isoformat() else "actif")
    contract = None
    if contract_type:
        contract = {"contract_type": contract_type, "position": values.get("position"), "start_date": start,
                    "end_date": values.get("contract_end"), "trial_end_date": trial_end, "salary_net": values.get("salary_net") or 0,
                    "status": "a_venir" if start and start > date.today().isoformat() else "actif", "link": row["link"]}
    return {"employee": employee, "contract": contract, "group": values.get("group"), "incomplete": row["incomplete"], "code_meta": row["code_meta"],
            "name": {key: row["name"].get(key) for key in ("original", "last_name", "first_name", "source", "confidence")}}


def _fingerprint(payload: dict[str, Any]) -> str:
    return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode()).hexdigest()[:32]


# ── Analyse complète (aperçu, corrections et revalidation à la confirmation) ───────────────
def _layout(session: CandidateImportSession, sheet_index: int) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook = session.workbook or {}
    sheets = workbook.get("sheets") or []
    if not 0 <= int(sheet_index) < len(sheets):
        raise HTTPException(status_code=422, detail="Feuille inconnue.")
    sheet = sheets[int(sheet_index)]
    layout = excel.sheet_layout(sheet, field_for_header)
    if not layout["importable"]:
        raise HTTPException(status_code=422, detail=f"Feuille « {sheet['name']} » non importable : {layout['reason']}.")
    return sheet, layout


def validate_mapping(layout: dict[str, Any], mapping: dict[str, Any], active_society: str | None) -> dict[int, str]:
    resolved = excel.validate_mapping(layout, mapping, fields=FIELD_BY_KEY, required=())
    fields = set(resolved.values())
    if "full_name" in fields and fields & {"last_name", "first_name"}:
        raise HTTPException(status_code=422, detail="Choisissez soit « Nom complet », soit les colonnes Nom et Prénom séparées.")
    if "full_name" not in fields and not {"last_name", "first_name"} <= fields:
        raise HTTPException(status_code=422, detail="Associez une colonne « Nom complet » ou les deux colonnes Nom et Prénom.")
    if "society" not in fields and not active_society:
        raise HTTPException(status_code=422, detail="Associez la colonne Société ou choisissez une société active.")
    return resolved


def analyze(db: Session, user: Any, session: CandidateImportSession, plan: dict[str, Any]) -> dict[str, Any]:
    active = (session.workbook or {}).get("society")
    sheet, layout = _layout(session, plan.get("sheet", 0))
    resolved = validate_mapping(layout, plan.get("mapping") or {}, active)
    ctx = _Context(db, user, active)
    overrides, accepts, excluded = plan.get("overrides") or {}, plan.get("accepts") or {}, set(plan.get("excluded") or [])
    rows = []
    for row_number, raw in layout["data_rows"]:
        cells = {key: (raw[index] if index < len(raw) else None) for index, key in resolved.items()}
        row = _analyze_row(ctx, row_number, cells, overrides.get(str(row_number)) or {}, set(accepts.get(str(row_number)) or []))
        row["excluded"] = str(row_number) in excluded
        row["accepted"] = sorted(accepts.get(str(row_number)) or [])
        rows.append(row)
    # Doublons dans le fichier : aucune des lignes concernées n'est créée tant qu'ils subsistent.
    active_rows = [row for row in rows if not row["excluded"]]
    seen_codes: dict[str, list[int]] = {}
    for row in active_rows:
        if row["values"].get("code"):
            seen_codes.setdefault(row["values"]["code"], []).append(row["row"])
    for kind, message in (("nin", "même NIN"), ("identity", "même nom et même téléphone")):
        groups: dict[str, list[int]] = {}
        for row in active_rows:
            if row["duplicate_keys"][kind]:
                groups.setdefault(row["duplicate_keys"][kind], []).append(row["row"])
        for numbers in groups.values():
            if len(numbers) > 1:
                for row in active_rows:
                    if row["row"] in numbers:
                        others = ", ".join(str(number) for number in numbers if number != row["row"])
                        row["issues"].append(_issue("blocked", "duplicate", f"Doublon dans le fichier ({message}) : ligne(s) {others}"))
    for code, numbers in seen_codes.items():
        if len(numbers) > 1:
            for row in active_rows:
                if row["row"] in numbers:
                    others = ", ".join(str(number) for number in numbers if number != row["row"])
                    row["issues"].append(_issue("blocked", "code", f"Matricule {code} présent plusieurs fois dans le fichier : ligne(s) {others}"))
    name_groups: dict[str, list[int]] = {}
    for row in active_rows:
        if row["duplicate_keys"]["name"]:
            name_groups.setdefault(row["duplicate_keys"]["name"], []).append(row["row"])
    for numbers in name_groups.values():
        if len(numbers) > 1:
            for row in active_rows:
                if row["row"] in numbers and not any(issue["field"] == "duplicate" and issue["level"] == "blocked" for issue in row["issues"]):
                    others = ", ".join(str(number) for number in numbers if number != row["row"])
                    issue = _issue("review", "duplicate", f"Même nom qu'à la ligne {others} du fichier", accept="duplicate")
                    if issue["accept"] not in set(accepts.get(str(row["row"])) or []):
                        row["issues"].append(issue)
    counts = {"total": len(rows), "ready": 0, "review": 0, "blocked": 0, "excluded": 0, "incomplete": 0,
              "empty_rows": layout.get("empty_rows", 0)}
    for row in rows:
        row["status"] = _status(row)
        counts[row["status"]] += 1
        if row["status"] == "ready":
            row["payload"] = _employee_payload(row)
            row["fingerprint"] = _fingerprint(row["payload"])
            counts["incomplete"] += bool(row["incomplete"])
    return {"sheet": {"index": int(plan.get("sheet", 0)), "name": sheet["name"], "header_row": layout["header_row"]},
            "mapping": {str(column): key for column, key in resolved.items()},
            "ignored_columns": [column["letter"] + (f" ({column['header']})" if column["header"] else "")
                                for column in layout["columns"] if column["index"] not in resolved],
            "society": active, "counts": counts, "rows": rows}


def public_analysis(analysis: dict[str, Any]) -> dict[str, Any]:
    rows = []
    for row in analysis["rows"]:
        values = dict(row["values"])
        rows.append({"row": row["row"], "status": row["status"], "excluded": row["excluded"], "issues": row["issues"], "accepted": row["accepted"],
                     "name": row["name"], "cells": row["cells"], "incomplete": row["incomplete"], "dropped": row["dropped"],
                     "values": {key: value for key, value in values.items() if value not in (None, "")},
                     "client_site": (row["link"] or {}).get("label"),
                     "code_preview": values.get("code") or "Automatique"})
    return {key: value for key, value in analysis.items() if key != "rows"} | {"rows": rows}


def _save_plan(session: CandidateImportSession, plan: dict[str, Any], analysis: dict[str, Any]) -> None:
    plan = {**plan, "ready": {str(row["row"]): row["fingerprint"] for row in analysis["rows"] if row["status"] == "ready"},
            "sheet_name": analysis["sheet"]["name"]}
    session.plan = plan
    session.status = "previewed"
    session.expires_at = datetime.utcnow() + timedelta(minutes=excel.SESSION_TTL_MINUTES)


# ── Sessions ────────────────────────────────────────────────────────────────────────────────
def upload(db: Session, user: Any, *, file_name: str, content: bytes, society: str | None, request: Any = None) -> dict[str, Any]:
    """Analyse le classeur (aucune fiche créée) et ouvre une session liée au compte, à son
    périmètre et à la société active choisie."""
    clean_name = re.sub(r"[\x00-\x1f]", "", str(file_name or "import.xlsx").replace("\\", "/").split("/")[-1])[:255] or "import.xlsx"
    scope_key, scope_label = excel.scope_fingerprint(user)
    try:
        workbook = excel.read_workbook(clean_name, content)
    except excel.ImportFileError as exc:
        append_audit(db, action="drh.employees.import.analyze", resource="employee_import", result="refused", user=user,
                     request=request, society=society or scope_label, new_state={"file_name": clean_name, "size": len(content), "reason": str(exc)})
        db.commit()
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    excel.purge_sessions(db)
    # Une seule analyse employés en attente par utilisateur : la précédente est abandonnée et vidée.
    db.execute(update(CandidateImportSession)
               .where(CandidateImportSession.user_id == user.id, CandidateImportSession.status.in_(("uploaded", "previewed")),
                      CandidateImportSession.public_id.startswith(PREFIX))
               .values(status="cancelled", workbook=None, plan=None))
    session = CandidateImportSession(
        public_id=PREFIX + secrets.token_hex(24), user_id=user.id, username=user.username, scope_key=scope_key, scope_label=scope_label,
        file_name=clean_name, file_sha256=hashlib.sha256(content).hexdigest(), file_size=len(content), status="uploaded",
        workbook={**workbook, "society": society}, expires_at=datetime.utcnow() + timedelta(minutes=excel.SESSION_TTL_MINUTES))
    db.add(session)
    append_audit(db, action="drh.employees.import.analyze", resource="employee_import", resource_id=session.public_id, result="success",
                 user=user, request=request, society=society or scope_label,
                 new_state={"file_name": clean_name, "size": len(content), "sheets": len(workbook["sheets"])})
    db.commit()
    db.refresh(session)
    return {"session_id": session.public_id, "file_name": session.file_name, "file_size": session.file_size, "society": society,
            "expires_at": session.expires_at.isoformat(), "limits": excel.limits(), **describe(workbook)}


def get_session(db: Session, user: Any, public_id: str, statuses: tuple[str, ...] | None = None) -> CandidateImportSession:
    return excel.get_session(db, user, public_id, statuses=statuses, purpose=PURPOSE)


def preview(db: Session, user: Any, session: CandidateImportSession, sheet_index: int, mapping: dict[str, Any]) -> dict[str, Any]:
    previous = session.plan or {}
    plan = {"sheet": int(sheet_index), "mapping": {str(key): value for key, value in (mapping or {}).items() if value}}
    if previous.get("sheet") == plan["sheet"] and previous.get("mapping") == plan["mapping"]:
        # Même feuille et même correspondance : les corrections déjà saisies sont conservées.
        plan.update({key: previous.get(key) or {} for key in ("overrides", "accepts")}, excluded=previous.get("excluded") or [])
    analysis = analyze(db, user, session, plan)
    _save_plan(session, plan, analysis)
    db.commit()
    return public_analysis(analysis)


def correct(db: Session, user: Any, session: CandidateImportSession, row_number: int, values: dict[str, Any],
            accept: list[str] | None, exclude: bool | None) -> dict[str, Any]:
    """Correction d'une ligne dans l'aperçu : enregistrée dans la session, puis tout le fichier
    est revérifié par le serveur (doublons internes compris)."""
    plan = dict(session.plan or {})
    _, layout = _layout(session, plan.get("sheet", 0))
    if row_number not in {number for number, _ in layout["data_rows"]}:
        raise HTTPException(status_code=422, detail="Ligne inconnue.")
    key = str(row_number)
    overrides = {**(plan.get("overrides") or {})}
    row_overrides = dict(overrides.get(key) or {})
    for field, value in (values or {}).items():
        if field not in FIELD_BY_KEY:
            raise HTTPException(status_code=422, detail="Correction invalide : champ inconnu.")
        if value is None:
            row_overrides.pop(field, None)                 # retour à la valeur du fichier
        else:
            text_value = " ".join(str(value).split())
            if len(text_value) > 300:
                raise HTTPException(status_code=422, detail="Correction trop longue.")
            row_overrides[field] = text_value
    overrides[key] = row_overrides
    plan["overrides"] = {k: v for k, v in overrides.items() if v}
    if accept is not None:
        accepts = {**(plan.get("accepts") or {})}
        accepts[key] = sorted({str(item)[:80] for item in accept})[:20]
        plan["accepts"] = {k: v for k, v in accepts.items() if v}
    if exclude is not None:
        excluded = set(plan.get("excluded") or [])
        (excluded.add if exclude else excluded.discard)(key)
        plan["excluded"] = sorted(excluded)
    analysis = analyze(db, user, session, plan)
    _save_plan(session, plan, analysis)
    db.commit()
    return public_analysis(analysis)


def cancel(db: Session, session: CandidateImportSession) -> None:
    excel.cancel(db, session)


# ── Confirmation ────────────────────────────────────────────────────────────────────────────
OUTCOME_LABELS = {"created": "Créé", "duplicate": "Doublon", "refused": "Refusé", "excluded": "Écarté", "error": "Erreur"}


def _outcome_for(row: dict[str, Any]) -> tuple[str, str]:
    if row["status"] == "excluded":
        return "excluded", "Ligne écartée dans l'aperçu"
    blocking = [issue for issue in row["issues"] if issue["level"] in ("blocked", "review")]
    duplicate = any(issue["field"] == "duplicate" or (issue["field"] == "code" and "déjà utilisé" in issue["message"])
                    or "plusieurs fois dans le fichier" in issue["message"] for issue in blocking if issue["level"] == "blocked")
    reason = " ; ".join(issue["message"] for issue in blocking) or "Ligne non prête"
    return ("duplicate" if duplicate else "refused"), reason


def _still_unique(db: Session, values: dict[str, Any]) -> str | None:
    """Recontrôle juste avant l'écriture, sous verrou : un employé créé entre-temps ailleurs."""
    if values.get("nin") and db.scalar(select(Employee.id).where(Employee.nin == values["nin"]).limit(1)) is not None:
        return "Employé créé entre-temps avec le même NIN"
    phone = _phone_key(values.get("phone"))
    if phone:
        rows = db.execute(select(Employee.phone, Employee.last_name, Employee.first_name).where(
            Employee.last_name == str(values["last_name"]).upper(), Employee.first_name == str(values["first_name"]).upper())).all()
        if any(_phone_key(row.phone) == phone for row in rows):
            return "Employé créé entre-temps avec le même nom et le même téléphone"
    return None


def confirm(db: Session, user: Any, public_id: str, *, request: Any = None) -> dict[str, Any]:
    """Crée les lignes prêtes, chacune dans sa propre transaction (import partiel). Idempotent :
    la session passe de « previewed » à « processing » par une mise à jour atomique ; un second
    appel ne crée rien et renvoie le résultat enregistré."""
    session = get_session(db, user, public_id)
    if session.status == "completed":
        return {**(session.result or {}), "already_processed": True}
    session = get_session(db, user, public_id, statuses=("previewed",))
    claimed = db.execute(update(CandidateImportSession)
                         .where(CandidateImportSession.id == session.id, CandidateImportSession.status == "previewed")
                         .values(status="processing", updated_at=datetime.utcnow())).rowcount
    db.commit()
    if claimed != 1:
        db.expire_all()
        current = db.get(CandidateImportSession, session.id)
        if current.status == "completed":
            return {**(current.result or {}), "already_processed": True}
        raise HTTPException(status_code=409, detail="Import déjà en cours de traitement.")
    db.expire_all()
    session = db.get(CandidateImportSession, session.id)
    plan = session.plan or {}
    try:
        analysis = analyze(db, user, session, plan)       # revalidation complète, données fraîches
    except Exception:
        session.status = "previewed"
        db.commit()
        raise
    approved = plan.get("ready") or {}
    meta = {"session_id": session.public_id, "file_name": session.file_name, "file_sha256": session.file_sha256,
            "sheet": analysis["sheet"]["name"], "society": analysis["society"]}
    report: dict[int, dict[str, Any]] = {}

    def snapshot(done: bool) -> dict[str, Any]:
        rows = [report.get(row["row"]) or {"row": row["row"], "outcome": "refused", "reason": "Traitement en attente ou interrompu"}
                for row in analysis["rows"]]
        counts = {"read": len(rows), "created": 0, "incomplete": 0, "duplicate": 0, "refused": 0, "excluded": 0, "error": 0,
                  "empty_rows": analysis["counts"]["empty_rows"]}
        for row in rows:
            counts[row["outcome"]] += 1
            counts["incomplete"] += bool(row["outcome"] == "created" and row.get("incomplete"))
        return {**meta, "counts": counts, "rows": rows, "completed_at": datetime.utcnow().isoformat() if done else None,
                "already_processed": False}

    for row in analysis["rows"]:
        if row["status"] != "ready":
            outcome, reason = _outcome_for(row)
            report[row["row"]] = {"row": row["row"], "outcome": outcome, "reason": reason, "original_name": row["name"].get("original"),
                                  "last_name": row["values"].get("last_name"), "first_name": row["values"].get("first_name"),
                                  "code": row["values"].get("code")}
        elif approved.get(str(row["row"])) != row["fingerprint"]:
            report[row["row"]] = {"row": row["row"], "outcome": "refused", "original_name": row["name"].get("original"),
                                  "last_name": row["values"].get("last_name"), "first_name": row["values"].get("first_name"),
                                  "reason": "La ligne a changé depuis la vérification : vérifiez-la à nouveau"}
    session.result = snapshot(False)
    db.commit()
    # Matricules saisis d'abord : un code attribué automatiquement ne peut pas prendre celui d'une ligne suivante.
    ready = [row for row in analysis["rows"] if row["row"] not in report]
    ready.sort(key=lambda row: (not row["values"].get("code"), row["row"]))
    for row in ready:
        payload = row["payload"]
        entry = {"row": row["row"], "original_name": payload["name"]["original"], "last_name": payload["employee"]["last_name"],
                 "first_name": payload["employee"]["first_name"], "incomplete": payload["incomplete"]}
        try:
            if db.get_bind().dialect.name == "postgresql":
                db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": _ADVISORY_LOCK_KEY})
            conflict = _still_unique(db, payload["employee"])
            if conflict:
                raise HTTPException(status_code=409, detail=conflict)
            employee_values = {key: value for key, value in payload["employee"].items() if value not in (None, "")}
            employee_values["extra"] = _extra(payload, meta, row, user)
            # Validation et mise en forme de la fiche DRH (majuscules, dates) ; matricule vide = série de la société.
            validated = EmployeeCreate(**{**employee_values, "code": employee_values.get("code") or "AUTO"}).model_dump()
            validated["code"] = employee_values.get("code") or ""
            contract = None
            if payload["contract"]:
                contract = {key: value for key, value in payload["contract"].items() if key != "link"}
                contract["position"] = validated["position"]
                for key in ("start_date", "end_date", "trial_end_date"):
                    contract[key] = date.fromisoformat(contract[key]) if contract[key] else None
                link = payload["contract"]["link"]
                if link:
                    contract.update(dc_client_sites.resolve_contract_selection(
                        db, dc_client_id=link["dc_client_id"], dc_site_id=link["dc_site_id"], society=payload["employee"]["society"]))
            employee = service.create_employee_direct(
                db, validated, strict_code=True, initial_contract=contract,
                actor=user, request=request, audit_action="drh.employees.import.created", commit=False,
                audit_state={"session_id": session.public_id, "file_sha256": session.file_sha256, "row": row["row"],
                             "original_full_name": payload["name"]["original"], "name_split": payload["name"],
                             "fiche_a_completer": bool(payload["incomplete"]), "incomplete_fields": payload["incomplete"],
                             "accepted": sorted(row["dropped"]), "ops_proposal": {"group": payload["group"]} if payload["group"] else None})
            entry.update(outcome="created", reason="Fiche à compléter : " + ", ".join(payload["incomplete"]) if payload["incomplete"] else "",
                         employee_id=employee.id, code=employee.code)
            report[row["row"]] = entry
            session.result = snapshot(False)
            db.commit()
        except HTTPException as exc:
            db.rollback()
            detail = str(exc.detail)
            duplicate = exc.status_code == 409 and ("Matricule" in detail or "même" in detail)
            entry.update(outcome="duplicate" if duplicate else "error", reason=detail, code=payload["employee"].get("code"))
            report[row["row"]] = entry
        except Exception:
            db.rollback()
            logger.exception("employee import row failed session=%s row=%s", session.public_id, row["row"])
            entry.update(outcome="error", reason="Erreur interne lors de l'enregistrement : aucune donnée conservée pour cette ligne",
                         code=payload["employee"].get("code"))
            report[row["row"]] = entry
        session = db.get(CandidateImportSession, session.id)
    session = db.get(CandidateImportSession, session.id)
    result = snapshot(True)
    session.status, session.result, session.completed_at = "completed", result, datetime.utcnow()
    session.workbook = None                                # contenu du fichier effacé dès la fin du traitement
    session.plan = {"sheet_name": analysis["sheet"]["name"]}
    append_audit(db, action="drh.employees.import", resource="employee_import", resource_id=session.public_id, result="success",
                 user=user, request=request, society=analysis["society"] or session.scope_label,
                 new_state={"file_name": session.file_name, "file_sha256": session.file_sha256, "sheet": analysis["sheet"]["name"],
                            **result["counts"]})
    db.commit()
    logger.info("employee import session=%s user=%s counts=%s", session.public_id, user.username, result["counts"])
    return result


def _extra(payload: dict[str, Any], meta: dict[str, Any], row: dict[str, Any], user: Any) -> dict[str, Any]:
    """Fiche à compléter (mêmes clés que le Recrutement groupé) et origine de l'import pour
    l'audit : valeur Excel originale du nom complet et séparation retenue. Le groupe n'est
    qu'une proposition pour OPS, jamais une affectation."""
    origin = {"source": "excel", "sessionId": meta["session_id"], "fileName": meta["file_name"], "fileSha256": meta["file_sha256"],
              "sheet": meta["sheet"], "row": row["row"], "importedAt": datetime.utcnow().isoformat(), "importedBy": user.username,
              "nomCompletOriginal": payload["name"]["original"],
              "separationNom": {"nom": payload["name"]["last_name"], "prenom": payload["name"]["first_name"],
                                "source": payload["name"]["source"], "confiance": payload["name"]["confidence"]}}
    if payload["group"]:
        origin["propositionOps"] = {"groupe": payload["group"], "statut": "à traiter par OPS"}
    if row["dropped"]:
        origin["champsEcartes"] = [label(field) for field in row["dropped"]]
    extra = {"ficheACompleter": bool(payload["incomplete"]), "recruitmentIncompleteFields": list(payload["incomplete"]), "importOrigin": origin}
    if payload["code_meta"]:
        # Matricule fourni par le fichier et validé : officiel, exclu de toute renumérotation.
        origin["matricule"] = payload["code_meta"]
        extra[code_policy.PROTECTION_KEY] = code_policy.protection_record(
            payload["code_meta"]["valeur"], source="import_excel", by=user.username, sessionId=meta["session_id"], ligne=row["row"],
            valeurExcel=payload["code_meta"]["valeurExcel"], formeProposee=payload["code_meta"]["formeProposee"], choix=payload["code_meta"]["choix"])
    return extra


def result_of(session: CandidateImportSession) -> dict[str, Any]:
    return {"status": session.status, **(session.result or {"rows": [], "counts": {}})}


# ── Modèle Excel et rapport ─────────────────────────────────────────────────────────────────
TEMPLATE_KEYS = ("code", "last_name", "first_name", "phone", "society", "client", "site", "position", "group", "recruit_date",
                 "contract_type", "contract_start", "contract_end", "nin", "birth_date", "birth_place", "family_status",
                 "children_count", "email", "address", "wilaya", "commune", "salary_net")


def build_template(db: Session, user: Any, society: str | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ctx = _Context(db, user, society)
    societies = [society] if society else sorted(ctx.allowed.values(), key=norm_key)
    positions = sorted(set(ctx.positions_for(society).values()) if society else {name for name, _ in ctx.positions}, key=norm_key)
    book = Workbook()
    sheet = book.active
    sheet.title = "Employés"
    lists = book.create_sheet("Listes")
    notes = book.create_sheet("Instructions")
    required_fill, header_fill = PatternFill("solid", fgColor="B45309"), PatternFill("solid", fgColor="043970")
    lists["A1"] = "LISTES DE RÉFÉRENCE — ne pas importer cette feuille"
    lists["A1"].font = Font(bold=True, color="B45309")
    ranges: dict[str, str] = {}
    columns = [("society", "Société", societies), ("position", "Poste", positions), ("group", "Groupe", list(GROUPS)),
               ("contract_type", "Type de contrat", list(CONTRACT_TYPES)),
               ("family_status", "Situation familiale", FIELD_BY_KEY["family_status"].choice_values())]
    for index, (key, title, values) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        lists.cell(row=3, column=index, value=title).font = Font(bold=True)
        for offset, value in enumerate(values, start=4):
            lists.cell(row=offset, column=index, value=excel.neutralize(value))
        lists.column_dimensions[letter].width = 28
        if values:
            ranges[key] = f"'Listes'!${letter}$4:${letter}${3 + len(values)}"
    required = {"last_name", "first_name"} | ({"society"} if not society else set())
    for index, key in enumerate(TEMPLATE_KEYS, start=1):
        item = FIELD_BY_KEY[key]
        letter = get_column_letter(index)
        cell = sheet.cell(row=1, column=index, value=item.label + (" *" if key in required else ""))
        cell.font, cell.fill = Font(bold=True, color="FFFFFF", size=10), required_fill if key in required else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if item.hint:
            cell.comment = Comment(item.hint, "ATLAS")
        sheet.column_dimensions[letter].width = item.width
        number_format = "@" if item.text_format else "DD/MM/YYYY" if item.kind == "date" else None
        if number_format:
            for row_index in range(2, excel.MAX_ROWS + 2):
                sheet.cell(row=row_index, column=index).number_format = number_format
        if key in ranges:
            validation = DataValidation(type="list", formula1=ranges[key], allow_blank=True, showErrorMessage=True,
                                        errorTitle="Valeur non autorisée", error="Choisissez une valeur de la liste.")
            sheet.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{excel.MAX_ROWS + 1}")
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    lines = [
        "INSTRUCTIONS — ne pas importer cette feuille",
        "",
        "1. Un employé par ligne dans la feuille « Employés », à partir de la ligne 2.",
        "2. Obligatoires : Nom et Prénom" + (", Société." if not society else f". Société : {society} (société active)."),
        "3. Nom et prénom dans une seule cellule acceptés (colonne « Nom complet ») : NOM en majuscules puis prénom(s), "
        "ex. BENALI Ahmed. ATLAS propose la séparation, vous la vérifiez avant l'import.",
        "4. Matricule vide : attribué par ATLAS selon la série de la société. Matricule déjà utilisé : ligne refusée.",
        "5. Société, client, site, poste et groupe doivent exister dans ATLAS : rien n'est créé automatiquement.",
        "6. Client et site sont rattachés au contrat initial, créé seulement si le type de contrat est renseigné.",
        "7. Le groupe (A, B, C ou D) est transmis à OPS comme proposition : aucune affectation n'est créée.",
        "8. Dates au format JJ/MM/AAAA. Téléphone et NIN : cellules au format Texte.",
        "9. Informations facultatives absentes : l'employé est créé avec le statut « Fiche à compléter ».",
        f"10. Limites : {excel.MAX_ROWS} lignes et {excel.MAX_FILE_BYTES // (1024 * 1024)} Mo par fichier, formats .xlsx ou .xls.",
    ]
    for index, line in enumerate(lines, start=1):
        notes.cell(row=index, column=1, value=line)
    notes["A1"].font = Font(bold=True, color="B45309")
    notes.column_dimensions["A"].width = 120
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def report_content(db: Session, user: Any, session: CandidateImportSession) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Après confirmation : résultat enregistré. Avant : état de la vérification en cours."""
    if session.status in ("completed", "processing", "failed") and session.result:
        return session.result["counts"], session.result["rows"]
    if session.status != "previewed":
        raise HTTPException(status_code=409, detail="Aucun rapport disponible pour cet import.")
    analysis = analyze(db, user, session, session.plan or {})
    labels = {"ready": "Prêt à importer", "review": "À compléter / à vérifier", "blocked": "Bloquant", "excluded": "Écarté"}
    rows = [{"row": row["row"], "outcome": labels[row["status"]], "original_name": row["name"].get("original"),
             "last_name": row["values"].get("last_name"), "first_name": row["values"].get("first_name"), "code": row["values"].get("code"),
             "reason": " ; ".join(issue["message"] for issue in row["issues"]) or ("Fiche à compléter : " + ", ".join(row["incomplete"]) if row["incomplete"] else "")}
            for row in analysis["rows"]]
    return {"read": analysis["counts"]["total"], **{labels[key]: analysis["counts"][key] for key in labels}}, rows


def build_report(session: CandidateImportSession, counts: dict[str, Any], rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Rapport d'import"
    names = {"read": "Lignes lues", "created": "Employés créés", "incomplete": "Dont fiches à compléter", "duplicate": "Doublons",
             "refused": "Lignes refusées", "excluded": "Lignes écartées", "error": "Erreurs", "empty_rows": "Lignes vides ignorées"}
    summary = [("Fichier", session.file_name), ("Feuille", (session.plan or {}).get("sheet_name") or (session.result or {}).get("sheet") or ""),
               ("Importé par", session.username or ""),
               ("Date", (session.completed_at or datetime.utcnow()).strftime("%d/%m/%Y %H:%M") + " UTC"),
               ("État", "Import terminé" if session.status == "completed" else "Vérification (aucun employé créé)")]
    summary += [(names.get(key, key), value) for key, value in counts.items()]
    for index, (title, value) in enumerate(summary, start=1):
        sheet.cell(row=index, column=1, value=title).font = Font(bold=True)
        cell = sheet.cell(row=index, column=2, value=excel.neutralize(value))
        if isinstance(value, str):
            cell.data_type = "s"
    header_row = len(summary) + 2
    headers = ("Ligne Excel", "Résultat", "Nom complet (valeur Excel)", "Nom", "Prénom", "Matricule", "Motif", "N° employé")
    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="043970")
    for offset, row in enumerate(rows, start=1):
        line = [row["row"], OUTCOME_LABELS.get(row["outcome"], row["outcome"]), row.get("original_name") or "", row.get("last_name") or "",
                row.get("first_name") or "", row.get("code") or "", row.get("reason") or "", row.get("employee_id") or ""]
        for index, value in enumerate(line, start=1):
            cell = sheet.cell(row=header_row + offset, column=index, value=excel.neutralize(value))
            if isinstance(value, str):
                cell.data_type = "s"                       # jamais de formule, même si la valeur commence par « = »
    for letter, width in zip("ABCDEFGH", (12, 16, 32, 22, 22, 14, 90, 12)):
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
